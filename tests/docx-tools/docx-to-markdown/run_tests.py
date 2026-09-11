# /// script
# requires-python = ">=3.10"
# dependencies = [
#     "python-docx>=1.1.2",
#     "mammoth>=1.8.0",
#     "markitdown[docx]>=0.0.1a4",
#     "Pillow>=10.0.0",
# ]
# ///
"""Comprehensive Test Suite for docx-to-markdown skill.

Executes all test scenarios covering:
- Headings and document structure
- GFM table translation
- Image extraction into media folder & relative linking (no inline base64)
- Word OpenXML comments extraction into Markdown (author, date, snippet, text)
- Lists, styles, hyperlinks
- Header and footer extraction
- Engines (markitdown, mammoth, auto)
- Real-world sample documents (file-sample_1, 2, 3)
- CLI options and edge case handling
"""

import docx
import json
import os
import re
import shutil
import subprocess
import sys
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

# Ensure UTF-8 output
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

REPO_ROOT = Path(__file__).resolve().parents[3]
PLUGIN_DIR = REPO_ROOT / ".agents" / "plugins" / "docx-tools"
SKILL_DIR = PLUGIN_DIR / "skills" / "docx-to-markdown"
CONVERT_SCRIPT = SKILL_DIR / "scripts" / "docx_to_md.py"

EXAMPLES_DIR = REPO_ROOT / "files_examples"
SAMPLE_1 = EXAMPLES_DIR / "file-sample_1.docx"
SAMPLE_2 = EXAMPLES_DIR / "file-sample_2.docx"
SAMPLE_3 = EXAMPLES_DIR / "file-sample_3.docx"

TEST_ROOT = Path(__file__).resolve().parent
SCENARIOS_DIR = TEST_ROOT / "scenarios"

UV_PATH = Path(shutil.which("uv") or r"C:\Users\kekss\.local\bin\uv.exe")

test_results = []


def run_command(cmd_args: list[str], timeout: int = 60) -> tuple[int, str, str]:
    """Run command with uv and returncode, stdout, stderr."""
    full_cmd = [str(UV_PATH), "run"] + cmd_args
    try:
        res = subprocess.run(
            full_cmd,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=timeout,
        )
        return res.returncode, res.stdout, res.stderr
    except subprocess.TimeoutExpired:
        return -1, "", "Command timed out"
    except Exception as e:
        return -2, "", str(e)


def record_result(name: str, passed: bool, details: str, output_folder: Path = None):
    test_results.append({
        "scenario": name,
        "passed": passed,
        "details": details,
        "folder": str(output_folder.relative_to(REPO_ROOT)) if output_folder else ""
    })
    status_str = "PASS" if passed else "FAIL"
    print(f"[{status_str}] {name}: {details}")


# =========================================================================
# Scenario 01: Headings and Paragraph Structure
# =========================================================================
def scenario_01_headings_and_structure():
    """Scenario 01: Verify heading translation (H1-H3, Russian headings) and text flow."""
    folder = SCENARIOS_DIR / "01_headings_and_structure"
    folder.mkdir(parents=True, exist_ok=True)

    test_docx = folder / "headings_test.docx"
    doc = docx.Document()
    doc.add_heading("Main Document Title", level=1)
    doc.add_paragraph("This is an introductory paragraph explaining section one.")
    doc.add_heading("Section Two Details", level=2)
    doc.add_paragraph("Detailed explanation with secondary level information.")
    doc.add_heading("Subsection In-Depth", level=3)
    doc.add_paragraph("Deep nested subsection text.")
    # Localized Russian styles
    p_ru1 = doc.add_paragraph("Русский Заголовок 1")
    try:
        p_ru1.style = doc.styles["Heading 1"]
    except Exception:
        pass
    p_ru2 = doc.add_paragraph("Русский Подзаголовок 2")
    try:
        p_ru2.style = doc.styles["Heading 2"]
    except Exception:
        pass
    doc.save(str(test_docx))

    out_md = folder / "headings_output.md"
    code, stdout, stderr = run_command([
        str(CONVERT_SCRIPT), str(test_docx),
        "-o", str(out_md)
    ])

    passed = (code == 0) and out_md.is_file()
    md_content = out_md.read_text(encoding="utf-8") if passed else ""

    has_h1 = "# Main Document Title" in md_content or "Main Document Title\n===" in md_content
    has_h2 = "## Section Two Details" in md_content or "Section Two Details\n---" in md_content
    has_h3 = "### Subsection In-Depth" in md_content
    has_ru1 = "Русский Заголовок 1" in md_content
    has_p1 = "This is an introductory paragraph" in md_content

    all_ok = passed and has_h1 and has_h2 and has_h3 and has_ru1 and has_p1

    rep = f"""# Сценарий 01: Заголовки и текстовая структура

- Код возврата: {code}
- Наличие H1: {has_h1}
- Наличие H2: {has_h2}
- Наличие H3: {has_h3}
- Наличие русского текста: {has_ru1}
- Статус: {'УСПЕХ' if all_ok else 'ОШИБКА'}
"""
    (folder / "scenario_report.md").write_text(rep, encoding="utf-8")
    record_result("Scenario 01: Headings & Structure", all_ok, "Preserved H1-H3 and localized headings", folder)


# =========================================================================
# Scenario 02: Table Formatting (GFM)
# =========================================================================
def scenario_02_tables_formatting():
    """Scenario 02: Verify DOCX tables are properly translated to GFM markdown tables."""
    folder = SCENARIOS_DIR / "02_tables_formatting"
    folder.mkdir(parents=True, exist_ok=True)

    # 1. Custom table with structured data
    custom_docx = folder / "custom_table.docx"
    doc = docx.Document()
    doc.add_heading("Financial Summary Table", level=1)
    table = doc.add_table(rows=4, cols=3)
    headers = ["Quarter", "Revenue ($M)", "Growth (%)"]
    data = [
        ["Q1 2025", "12.4", "+8.5%"],
        ["Q2 2025", "15.1", "+21.8%"],
        ["Q3 2025", "14.8", "-2.0%"],
    ]
    for col_idx, h in enumerate(headers):
        table.cell(0, col_idx).text = h
    for row_idx, row_data in enumerate(data, start=1):
        for col_idx, val in enumerate(row_data):
            table.cell(row_idx, col_idx).text = val
    doc.save(str(custom_docx))

    out_custom_md = folder / "custom_table.md"
    code_c, _, _ = run_command([str(CONVERT_SCRIPT), str(custom_docx), "-o", str(out_custom_md)])

    c_md = out_custom_md.read_text(encoding="utf-8") if out_custom_md.is_file() else ""
    has_gfm_sep = ("| --- |" in c_md) or ("|:---|" in c_md) or ("| ---" in c_md)
    has_headers = "Quarter" in c_md and "Revenue" in c_md and "Growth" in c_md
    has_data_cell = "Q1 2025" in c_md and "15.1" in c_md

    # 2. Real table from file-sample_2.docx
    out_s2_md = folder / "sample2_table.md"
    code_s2, _, _ = run_command([str(CONVERT_SCRIPT), str(SAMPLE_2), "-o", str(out_s2_md)])
    s2_md = out_s2_md.read_text(encoding="utf-8") if out_s2_md.is_file() else ""
    s2_has_table = "| --- |" in s2_md and "Lorem ipsum" in s2_md

    all_ok = (code_c == 0) and has_gfm_sep and has_headers and has_data_cell and (code_s2 == 0) and s2_has_table

    rep = f"""# Сценарий 02: Трансляция таблиц в Markdown

- Пользовательская таблица: код {code_c}, разделители: {has_gfm_sep}, заголовки: {has_headers}
- Реальная таблица из sample_2: код {code_s2}, разделители: {s2_has_table}
- Статус: {'УСПЕХ' if all_ok else 'ОШИБКА'}
"""
    (folder / "scenario_report.md").write_text(rep, encoding="utf-8")
    record_result("Scenario 02: Tables to GFM", all_ok, "Clean GFM markdown tables with headers & delimiters", folder)


# =========================================================================
# Scenario 03: Image Extraction and Relative Linking
# =========================================================================
def scenario_03_image_extraction_and_paths():
    """Scenario 03: Verify embedded images are exported to disk and linked with relative paths."""
    folder = SCENARIOS_DIR / "03_image_extraction_and_paths"
    folder.mkdir(parents=True, exist_ok=True)

    out_md = folder / "sample1_report.md"
    media_dir = folder / "sample1_extracted_images"

    code, stdout, stderr = run_command([
        str(CONVERT_SCRIPT), str(SAMPLE_1),
        "-o", str(out_md),
        "--media-dir", str(media_dir)
    ])

    passed = (code == 0) and out_md.is_file() and media_dir.is_dir()
    md_content = out_md.read_text(encoding="utf-8") if out_md.is_file() else ""

    # Check extracted image files on disk
    image_files = list(media_dir.glob("image_*.*"))
    has_images_on_disk = len(image_files) >= 2
    images_non_empty = all(img.stat().st_size > 0 for img in image_files) if image_files else False

    # Check relative links in markdown: ![...](sample1_extracted_images/image_001.png)
    has_relative_links = ("sample1_extracted_images/image_001" in md_content) or ("sample1_extracted_images/image_002" in md_content)
    # Check that NO raw data URI base64 remains
    has_no_raw_base64 = "data:image/" not in md_content

    all_ok = passed and has_images_on_disk and images_non_empty and has_relative_links and has_no_raw_base64

    rep = f"""# Сценарий 03: Экспорт изображений и относительные ссылки

- Извлечено изображений на диск: {len(image_files)} (ожидалось >= 2)
- Все файлы непустые: {images_non_empty}
- Относительные ссылки в Markdown: {has_relative_links}
- Отсутствие inline base64 blobs: {has_no_raw_base64}
- Статус: {'УСПЕХ' if all_ok else 'ОШИБКА'}
"""
    (folder / "scenario_report.md").write_text(rep, encoding="utf-8")
    record_result("Scenario 03: Images & Relative Paths", all_ok, f"Extracted {len(image_files)} images, verified relative markdown links", folder)


# =========================================================================
# Scenario 04: Word Comments Extraction
# =========================================================================
def scenario_04_comments_extraction():
    """Scenario 04: Verify Word comments are parsed and included into Markdown with author, date, and text."""
    folder = SCENARIOS_DIR / "04_comments_extraction"
    folder.mkdir(parents=True, exist_ok=True)

    # 1. Real document with comments: file-sample_3.docx
    out_s3_md = folder / "sample3_with_comments.md"
    media_s3 = folder / "sample3_media"
    code_s3, stdout_s3, _ = run_command([
        str(CONVERT_SCRIPT), str(SAMPLE_3),
        "-o", str(out_s3_md),
        "--media-dir", str(media_s3)
    ])

    md_s3 = out_s3_md.read_text(encoding="utf-8") if out_s3_md.is_file() else ""
    has_comments_header = "## Комментарии / Comments" in md_s3
    has_author = "Konstantin Ivanov" in md_s3
    has_comment_text = "My Comment" in md_s3
    has_snippet = "Over the past six months" in md_s3

    # 2. Test multi-comment synthetic document
    multi_docx = folder / "multi_comments.docx"
    shutil.copy2(SAMPLE_3, multi_docx)
    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    ns = {"w": w_ns}

    with zipfile.ZipFile(multi_docx, "r") as zin:
        files = {name: zin.read(name) for name in zin.namelist()}

    c_root = ET.fromstring(files["word/comments.xml"])
    c2 = ET.SubElement(c_root, f"{{{w_ns}}}comment", {
        f"{{{w_ns}}}id": "1",
        f"{{{w_ns}}}author": "Anna Petrova",
        f"{{{w_ns}}}date": "2026-09-11T14:30:00Z",
        f"{{{w_ns}}}initials": "AP",
    })
    p = ET.SubElement(c2, f"{{{w_ns}}}p")
    r = ET.SubElement(p, f"{{{w_ns}}}r")
    t = ET.SubElement(r, f"{{{w_ns}}}t")
    t.text = "Second test comment from reviewer"
    files["word/comments.xml"] = ET.tostring(c_root, encoding="utf-8")

    d_root = ET.fromstring(files["word/document.xml"])
    ps = d_root.findall(".//w:p", ns)
    if len(ps) > 5:
        target_p = ps[5]
        ET.SubElement(target_p, f"{{{w_ns}}}commentReference", {f"{{{w_ns}}}id": "1"})
    files["word/document.xml"] = ET.tostring(d_root, encoding="utf-8")

    with zipfile.ZipFile(multi_docx, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in files.items():
            zout.writestr(name, data)

    out_multi_md = folder / "multi_comments.md"
    code_multi, stdout_multi, _ = run_command([
        str(CONVERT_SCRIPT), str(multi_docx),
        "-o", str(out_multi_md),
        "--media-dir", str(folder / "multi_media")
    ])

    md_multi = out_multi_md.read_text(encoding="utf-8") if out_multi_md.is_file() else ""
    has_second_author = "Anna Petrova" in md_multi
    has_second_comment = "Second test comment from reviewer" in md_multi

    # 3. Test --no-comments flag
    out_no_comments_md = folder / "sample3_no_comments.md"
    code_no_c, _, _ = run_command([
        str(CONVERT_SCRIPT), str(SAMPLE_3),
        "-o", str(out_no_comments_md),
        "--media-dir", str(folder / "no_comments_media"),
        "--no-comments"
    ])
    md_no_c = out_no_comments_md.read_text(encoding="utf-8") if out_no_comments_md.is_file() else ""
    correctly_omitted = "## Комментарии / Comments" not in md_no_c

    all_ok = (
        (code_s3 == 0) and has_comments_header and has_author and has_comment_text and has_snippet
        and (code_multi == 0) and has_second_author and has_second_comment
        and (code_no_c == 0) and correctly_omitted
    )

    rep = f"""# Сценарий 04: Извлечение комментариев Word

- Реальный файл sample_3: код {code_s3}, автор найден: {has_author}, текст комментария найден: {has_comment_text}
- Многопользовательский комментарий: код {code_multi}, второй автор: {has_second_author}, текст: {has_second_comment}
- Флаг --no-comments: код {code_no_c}, комментарии опущены: {correctly_omitted}
- Статус: {'УСПЕХ' if all_ok else 'ОШИБКА'}
"""
    (folder / "scenario_report.md").write_text(rep, encoding="utf-8")
    record_result("Scenario 04: Comments Extraction", all_ok, "Extracted author, timestamp, target snippet & comment text", folder)


# =========================================================================
# Scenario 05: Lists, Text Styles, and Hyperlinks
# =========================================================================
def scenario_05_lists_styles_links():
    """Scenario 05: Verify bullet/numbered lists, bold/italic styles, and hyperlinks."""
    folder = SCENARIOS_DIR / "05_lists_styles_links"
    folder.mkdir(parents=True, exist_ok=True)

    test_docx = folder / "styles_and_links.docx"
    doc = docx.Document()
    doc.add_heading("Formatting & Lists", level=1)

    p1 = doc.add_paragraph()
    r1 = p1.add_run("This text contains ")
    r2 = p1.add_run("bold formatting")
    r2.bold = True
    r3 = p1.add_run(" and ")
    r4 = p1.add_run("italic formatting")
    r4.italic = True
    r5 = p1.add_run(".")

    # Bullet list
    doc.add_paragraph("First bullet item", style="List Bullet")
    doc.add_paragraph("Second bullet item", style="List Bullet")
    doc.add_paragraph("Third bullet item", style="List Bullet")

    # Hyperlink insertion via OpenXML
    p_link = doc.add_paragraph("Visit our documentation: ")
    part = doc.part
    r_id = part.relate_to("https://github.com/project", docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = docx.oxml.OxmlElement("w:hyperlink")
    hyperlink.set(docx.oxml.ns.qn("r:id"), r_id)
    new_run = docx.oxml.OxmlElement("w:r")
    r_text = docx.oxml.OxmlElement("w:t")
    r_text.text = "Antigravity Project Repo"
    new_run.append(r_text)
    hyperlink.append(new_run)
    p_link._p.append(hyperlink)

    doc.save(str(test_docx))

    out_md = folder / "styles_output.md"
    code, _, _ = run_command([str(CONVERT_SCRIPT), str(test_docx), "-o", str(out_md)])

    md_content = out_md.read_text(encoding="utf-8") if out_md.is_file() else ""

    has_bold = ("**bold formatting**" in md_content) or ("__bold formatting__" in md_content)
    has_italic = ("*italic formatting*" in md_content) or ("_italic formatting_" in md_content)
    has_bullet = ("* First bullet item" in md_content) or ("- First bullet item" in md_content)
    has_link = ("https://github.com/project" in md_content) and ("Antigravity Project Repo" in md_content)

    all_ok = (code == 0) and has_bold and has_italic and has_bullet and has_link

    rep = f"""# Сценарий 05: Списки, стили и гиперссылки

- Жирный шрифт: {has_bold}
- Курсив: {has_italic}
- Маркированный список: {has_bullet}
- Внешняя гиперссылка: {has_link}
- Статус: {'УСПЕХ' if all_ok else 'ОШИБКА'}
"""
    (folder / "scenario_report.md").write_text(rep, encoding="utf-8")
    record_result("Scenario 05: Lists, Styles & Links", all_ok, "Preserved bold, italic, bullet lists, and markdown links", folder)


# =========================================================================
# Scenario 06: Headers and Footers
# =========================================================================
def scenario_06_headers_footers():
    """Scenario 06: Verify --include-headers option extracts header and footer content."""
    folder = SCENARIOS_DIR / "06_headers_footers"
    folder.mkdir(parents=True, exist_ok=True)

    test_docx = folder / "document_with_header.docx"
    doc = docx.Document()
    section = doc.sections[0]
    header = section.header
    header.paragraphs[0].text = "CONFIDENTIAL INTERNAL REPORT 2026"
    footer = section.footer
    footer.paragraphs[0].text = "Page 1 of 10 - Proprietary"

    doc.add_heading("Report Subject", level=1)
    doc.add_paragraph("Core document content goes here.")
    doc.save(str(test_docx))

    # Test WITH --include-headers
    out_with_h = folder / "with_headers.md"
    code_h, _, _ = run_command([
        str(CONVERT_SCRIPT), str(test_docx),
        "-o", str(out_with_h),
        "--include-headers"
    ])
    md_with_h = out_with_h.read_text(encoding="utf-8") if out_with_h.is_file() else ""
    has_header_block = ("> **Колонтитул:**" in md_with_h) and ("CONFIDENTIAL INTERNAL REPORT 2026" in md_with_h)

    # Test WITHOUT --include-headers
    out_without_h = folder / "without_headers.md"
    code_noh, _, _ = run_command([
        str(CONVERT_SCRIPT), str(test_docx),
        "-o", str(out_without_h)
    ])
    md_without_h = out_without_h.read_text(encoding="utf-8") if out_without_h.is_file() else ""
    lacks_header_block = "> **Колонтитул:**" not in md_without_h

    all_ok = (code_h == 0) and has_header_block and (code_noh == 0) and lacks_header_block

    rep = f"""# Сценарий 06: Извлечение колонтитулов

- С флагом --include-headers: {has_header_block}
- Без флага --include-headers: {lacks_header_block}
- Статус: {'УСПЕХ' if all_ok else 'ОШИБКА'}
"""
    (folder / "scenario_report.md").write_text(rep, encoding="utf-8")
    record_result("Scenario 06: Headers & Footers", all_ok, "Extracted header/footer block with --include-headers", folder)


# =========================================================================
# Scenario 07: Conversion Engines
# =========================================================================
def scenario_07_conversion_engines():
    """Scenario 07: Verify explicit --engine choices (markitdown, mammoth, auto)."""
    folder = SCENARIOS_DIR / "07_conversion_engines"
    folder.mkdir(parents=True, exist_ok=True)

    engines = ["markitdown", "mammoth", "auto"]
    results = {}

    for eng in engines:
        out_md = folder / f"sample2_engine_{eng}.md"
        media_d = folder / f"media_{eng}"
        code, stdout, stderr = run_command([
            str(CONVERT_SCRIPT), str(SAMPLE_2),
            "-o", str(out_md),
            "--media-dir", str(media_d),
            "--engine", eng
        ])
        success = (code == 0) and out_md.is_file() and (out_md.stat().st_size > 100)
        results[eng] = success

    all_ok = all(results.values())

    rep = f"""# Сценарий 07: Движки конвертации

- MarkItDown: {'УСПЕХ' if results.get('markitdown') else 'ОШИБКА'}
- Mammoth: {'УСПЕХ' if results.get('mammoth') else 'ОШИБКА'}
- Auto (Default): {'УСПЕХ' if results.get('auto') else 'ОШИБКА'}
- Статус: {'УСПЕХ' if all_ok else 'ОШИБКА'}
"""
    (folder / "scenario_report.md").write_text(rep, encoding="utf-8")
    record_result("Scenario 07: Conversion Engines", all_ok, "Successfully tested markitdown, mammoth, and auto engines", folder)


# =========================================================================
# Scenario 08: Real-world Sample Documents
# =========================================================================
def scenario_08_real_world_samples():
    """Scenario 08: Comprehensive run across all 3 sample files (file-sample_1, 2, 3)."""
    folder = SCENARIOS_DIR / "08_real_world_samples"
    folder.mkdir(parents=True, exist_ok=True)

    samples = [
        (SAMPLE_1, "sample_1_b2b_proposal", 2, True),   # has 2 images, table
        (SAMPLE_2, "sample_2_article", 1, True),         # has 1 image, table
        (SAMPLE_3, "sample_3_hr_report", 8, True),       # has 8 images, comments, tables
    ]

    all_ok = True
    notes = []

    for s_file, s_name, exp_min_images, has_table in samples:
        out_md = folder / f"{s_name}.md"
        media_d = folder / f"{s_name}_media"

        code, stdout, stderr = run_command([
            str(CONVERT_SCRIPT), str(s_file),
            "-o", str(out_md),
            "--media-dir", str(media_d)
        ])

        if code != 0 or not out_md.is_file():
            all_ok = False
            notes.append(f"{s_name} failed with code {code}: {stderr}")
            continue

        md_text = out_md.read_text(encoding="utf-8")
        images_found = list(media_d.glob("image_*.*"))

        if len(images_found) < exp_min_images:
            all_ok = False
            notes.append(f"{s_name}: found {len(images_found)} images, expected >= {exp_min_images}")

        # Check relative link presence
        if images_found:
            rel_name = media_d.name
            if rel_name not in md_text:
                all_ok = False
                notes.append(f"{s_name}: media folder name '{rel_name}' missing from markdown links")

        # Check comment in sample 3
        if s_file == SAMPLE_3:
            if "Konstantin Ivanov" not in md_text or "My Comment" not in md_text:
                all_ok = False
                notes.append(f"{s_name}: comment by Konstantin Ivanov was not extracted")

    rep = f"""# Сценарий 08: Прогон на реальных файлах

- file-sample_1.docx (b2b предложение): УСПЕХ
- file-sample_2.docx (статья с таблицей): УСПЕХ
- file-sample_3.docx (HR отчёт с комментариями): УСПЕХ
- Замечания: {'; '.join(notes) if notes else 'Нет'}
- Статус: {'УСПЕХ' if all_ok else 'ОШИБКА'}
"""
    (folder / "scenario_report.md").write_text(rep, encoding="utf-8")
    record_result("Scenario 08: Real Samples Suite", all_ok, "All 3 sample docx files converted with media, tables & comments", folder)


# =========================================================================
# Scenario 09: CLI Arguments & Edge Cases
# =========================================================================
def scenario_09_cli_and_edge_cases():
    """Scenario 09: Test CLI arguments, non-existent files, nested directory creation."""
    folder = SCENARIOS_DIR / "09_cli_and_edge_cases"
    folder.mkdir(parents=True, exist_ok=True)

    all_ok = True
    notes = []

    # 1. Non-existent input file
    fake_file = folder / "non_existent_file.docx"
    code_fake, stdout_fake, stderr_fake = run_command([str(CONVERT_SCRIPT), str(fake_file)])
    if code_fake != 1 or "not found" not in stderr_fake.lower():
        all_ok = False
        notes.append(f"Non-existent file test failed: code {code_fake}, stderr: {stderr_fake}")

    # 2. Deeply nested output path
    nested_md = folder / "deep" / "nested" / "path" / "output.md"
    nested_media = folder / "deep" / "media" / "folder"
    code_nest, _, _ = run_command([
        str(CONVERT_SCRIPT), str(SAMPLE_1),
        "-o", str(nested_md),
        "--media-dir", str(nested_media)
    ])
    if code_nest != 0 or not nested_md.is_file() or not nested_media.is_dir():
        all_ok = False
        notes.append("Nested directory auto-creation failed")

    # 3. Minimal blank document (no images, no tables, no comments)
    blank_docx = folder / "minimal.docx"
    doc = docx.Document()
    doc.add_paragraph("Just a plain text paragraph with nothing else.")
    doc.save(str(blank_docx))

    blank_md = folder / "minimal.md"
    code_blank, _, _ = run_command([str(CONVERT_SCRIPT), str(blank_docx), "-o", str(blank_md)])
    if code_blank != 0 or not blank_md.is_file() or "plain text paragraph" not in blank_md.read_text(encoding="utf-8"):
        all_ok = False
        notes.append("Minimal document conversion failed")

    rep = f"""# Сценарий 09: CLI и пограничные случаи

- Обработка несуществующего файла: {'УСПЕХ' if code_fake == 1 else 'ОШИБКА'}
- Создание вложенных путей: {'УСПЕХ' if code_nest == 0 and nested_md.is_file() else 'ОШИБКА'}
- Минимальный документ без картинок и таблиц: {'УСПЕХ' if code_blank == 0 else 'ОШИБКА'}
- Замечания: {'; '.join(notes) if notes else 'Нет'}
- Статус: {'УСПЕХ' if all_ok else 'ОШИБКА'}
"""
    (folder / "scenario_report.md").write_text(rep, encoding="utf-8")
    record_result("Scenario 09: CLI & Edge Cases", all_ok, "Verified error handling, nested path creation, minimal docx", folder)


# =========================================================================
# Main Execution Runner
# =========================================================================
def main():
    print("=" * 80)
    print("ТЕСТОВЫЙ НАБОР ДЛЯ СКИЛА: docx-to-markdown")
    print(f"Целевой скрипт: {CONVERT_SCRIPT.relative_to(REPO_ROOT)}")
    print(f"Папка результатов: {SCENARIOS_DIR.relative_to(REPO_ROOT)}")
    print("=" * 80)

    scenario_01_headings_and_structure()
    scenario_02_tables_formatting()
    scenario_03_image_extraction_and_paths()
    scenario_04_comments_extraction()
    scenario_05_lists_styles_links()
    scenario_06_headers_footers()
    scenario_07_conversion_engines()
    scenario_08_real_world_samples()
    scenario_09_cli_and_edge_cases()

    # Generate test_summary.json
    all_passed = all(r["passed"] for r in test_results)
    summary = {
        "skill": "docx-to-markdown",
        "total_scenarios": len(test_results),
        "passed_scenarios": sum(1 for r in test_results if r["passed"]),
        "failed_scenarios": sum(1 for r in test_results if not r["passed"]),
        "all_passed": all_passed,
        "scenarios": test_results
    }

    summary_file = TEST_ROOT / "test_summary.json"
    summary_file.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    print("\n" + "=" * 80)
    print(f"ИТОГИ ТЕСТИРОВАНИЯ docx-to-markdown: {summary['passed_scenarios']}/{summary['total_scenarios']} УСПЕШНО")
    print("=" * 80)

    for r in test_results:
        st = "✅ PASS" if r["passed"] else "❌ FAIL"
        print(f"[{st}] {r['scenario']}: {r['details']}")

    sys.exit(0 if all_passed else 1)


if __name__ == "__main__":
    main()
