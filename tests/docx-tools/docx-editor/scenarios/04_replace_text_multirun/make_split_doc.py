
import docx
doc = docx.Document()
p = doc.add_paragraph()
r1 = p.add_run("Project ")
r1.bold = False
r2 = p.add_run("Antigravity")
r2.bold = True
r2.italic = False
r3 = p.add_run(" IDE Platform")
r3.italic = True
doc.save(r"c:/Users/kekss/Project/Antigravity_IDE_configs/test_docx_editor_suite/scenarios/04_replace_text_multirun/test_split.docx")
