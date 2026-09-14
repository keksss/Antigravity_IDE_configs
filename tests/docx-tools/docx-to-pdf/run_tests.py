# /// script
# requires-python = ">=3.10"
# dependencies = [
#     "python-docx>=1.1.2",
#     "pymupdf>=1.24.0",
#     "Pillow>=10.0.0",
#     "pywin32; sys_platform == 'win32'",
#     "docx2pdf>=0.1.8; sys_platform == 'win32' or sys_platform == 'darwin'",
# ]
# ///
"""Comprehensive Test Suite for docx-to-pdf skill.

Executes all 9 test scenarios covering:
- Scenario 01: Single Document Conversion with Word COM (--engine word)
- Scenario 02: Engine Selection, Auto-detection & LibreOffice Fallback
- Scenario 03: Custom Output Path and Nested Destination Directory Creation
- Scenario 04: Batch Directory Conversion
- Scenario 05: Overwrite Protection and --force Flag
- Scenario 06: Filtering Word Lock Files (~$*.docx)
- Scenario 07: Error Handling and Edge Cases (Corrupted, 0-byte, Missing Files)
- Scenario 08: Real-World Samples Suite (file-sample_1, 2, 3)
- Scenario 09: PDF Quality & Content Inspection (PyMuPDF Text, Tables, Images)
"""

import json
import os
import platform
import shutil
import subprocess
import sys
import time
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import docx
from docx.shared import Inches, Pt, RGBColor
import fitz  # PyMuPDF
from PIL import Image, ImageDraw

# Ensure UTF-8 output on Windows consoles
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

REPO_ROOT = Path(__file__).resolve().parents[3]
PLUGIN_DIR = REPO_ROOT / ".agents" / "plugins" / "docx-tools"
SKILL_DIR = PLUGIN_DIR / "skills" / "docx-to-pdf"
CONVERT_SCRIPT = SKILL_DIR / "scripts" / "docx_to_pdf.py"

EXAMPLES_DIR = REPO_ROOT / "files_examples"
SAMPLE_1 = EXAMPLES_DIR / "file-sample_1.docx"
SAMPLE_2 = EXAMPLES_DIR / "file-sample_2.docx"
SAMPLE_3 = EXAMPLES_DIR / "file-sample_3.docx"

TEST_ROOT = Path(__file__).resolve().parent
SCENARIOS_DIR = TEST_ROOT / "scenarios"

def get_uv_path() -> Optional[Path]:
    for candidate in [
        shutil.which("uv"),
        REPO_ROOT / "uv.exe",
        Path.home() / ".local" / "bin" / "uv.exe",
        Path.home() / ".cargo" / "bin" / "uv.exe",
    ]:
        if candidate:
            try:
                p = Path(candidate)
                if p.is_file():
                    return p
            except Exception:
                pass
    return None

UV_PATH = get_uv_path()


test_results: List[Dict[str, Any]] = []


def run_command(cmd_args: List[str], timeout: int = 90) -> Tuple[int, str, str]:
    """Execute command via uv run or directly with python."""
    if UV_PATH and UV_PATH.is_file():
        full_cmd = [str(UV_PATH), "run"] + cmd_args
    else:
        full_cmd = [sys.executable] + cmd_args

    try:
        res = subprocess.run(
            full_cmd,
            cwd=str(REPO_ROOT),
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=timeout,
        )
        return res.returncode, res.stdout, res.stderr
    except subprocess.TimeoutExpired:
        return -1, "", f"Command timed out after {timeout} seconds"
    except Exception as e:
        return -2, "", str(e)


def record_result(name: str, passed: bool, details: str, output_folder: Optional[Path] = None):
    test_results.append({
        "scenario": name,
        "passed": passed,
        "details": details,
        "folder": str(output_folder.relative_to(REPO_ROOT)) if output_folder else "",
    })
    status_str = "PASS" if passed else "FAIL"
    print(f"[{status_str}] {name}: {details}")


def is_valid_pdf_file(pdf_path: Path) -> Tuple[bool, str]:
    """Check if file exists, is non-empty, and has valid PDF header."""
    if not pdf_path.is_file():
        return False, "File does not exist"
    if pdf_path.stat().st_size == 0:
        return False, "File is 0 bytes"
    try:
        with open(pdf_path, "rb") as f:
            header = f.read(5)
            if not header.startswith(b"%PDF-"):
                return False, f"Invalid PDF header: {header!r}"
    except Exception as e:
        return False, f"Error reading PDF header: {e}"
    return True, "Valid PDF"


def create_sample_image(img_path: Path, width: int = 400, height: int = 200, label: str = "Test Image"):
    """Generate a sample PNG image with shapes and text for embedding."""
    img = Image.new("RGB", (width, height), color=(235, 240, 250))
    draw = ImageDraw.Draw(img)
    # Draw decorative rectangle and lines
    draw.rectangle([10, 10, width - 10, height - 10], outline=(40, 80, 160), width=3)
    draw.line([20, height // 2, width - 20, height // 2], fill=(180, 40, 40), width=2)
    draw.text((25, 25), label, fill=(10, 30, 80))
    img.save(img_path)


# =========================================================================
# Scenario 01: Single Document Conversion with Word COM (--engine word)
# =========================================================================
def scenario_01_single_file_word():
    """Scenario 01: Convert rich .docx with Word COM and verify output structure."""
    folder = SCENARIOS_DIR / "01_single_file_word"
    folder.mkdir(parents=True, exist_ok=True)

    test_docx = folder / "single_test.docx"
    doc = docx.Document()
    doc.add_heading("Отчёт по тестированию Word COM", level=0)

    p1 = doc.add_paragraph("Тестовый документ для проверки нативной конвертации через ")
    run_bold = p1.add_run("Microsoft Word COM Automation.")
    run_bold.bold = True

    doc.add_heading("1. Табличные данные", level=1)
    table = doc.add_table(rows=3, cols=3)
    headers = ["Параметр", "Значение", "Статус"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
    rows_data = [
        ["Движок", "Word COM", "Активен"],
        ["Рендеринг", "Нативный", "100% точность"],
    ]
    for r_idx, row in enumerate(rows_data, 1):
        for c_idx, val in enumerate(row):
            table.cell(r_idx, c_idx).text = val

    img_path = folder / "test_banner.png"
    create_sample_image(img_path, 300, 120, "Word COM Banner")
    doc.add_heading("2. Графические объекты", level=1)
    doc.add_picture(str(img_path), width=Inches(3.5))

    doc.save(str(test_docx))

    output_pdf = folder / "single_test.pdf"
    if output_pdf.exists():
        output_pdf.unlink()

    code, stdout, stderr = run_command([
        str(CONVERT_SCRIPT),
        str(test_docx),
        "-o",
        str(output_pdf),
        "--engine",
        "word",
        "--force",
    ])

    all_ok = True
    notes = []

    if code != 0:
        all_ok = False
        notes.append(f"Exit code {code}, stderr: {stderr.strip()}")

    valid, reason = is_valid_pdf_file(output_pdf)
    if not valid:
        all_ok = False
        notes.append(f"PDF validation failed: {reason}")
    else:
        # Inspect with PyMuPDF
        try:
            pdf_doc = fitz.open(str(output_pdf))
            if len(pdf_doc) < 1:
                all_ok = False
                notes.append("PDF has 0 pages")
            text = "".join(page.get_text() for page in pdf_doc)
            if "Word COM" not in text:
                all_ok = False
                notes.append("Expected text 'Word COM' missing in converted PDF")
            pdf_doc.close()
        except Exception as e:
            all_ok = False
            notes.append(f"PyMuPDF inspection error: {e}")

    report = f"""# Сценарий 01: Одиночная конвертация через Word COM

- Исходный документ: `{test_docx.name}`
- Результат PDF: `{output_pdf.name}`
- Код возврата CLI: {code}
- Размер PDF: {output_pdf.stat().st_size if output_pdf.exists() else 0} байт
- Замечания: {'; '.join(notes) if notes else 'Нет'}
- Статус: {'УСПЕХ' if all_ok else 'ОШИБКА'}
"""
    (folder / "scenario_report.md").write_text(report, encoding="utf-8")
    record_result("Scenario 01: Single File Word COM", all_ok, "Native Word COM conversion and PDF structure verification", folder)


# =========================================================================
# Scenario 02: Engine Selection, Auto-detection & LibreOffice Fallback
# =========================================================================
def scenario_02_engine_selection():
    """Scenario 02: Test --engine auto, detection helpers, and invalid engine errors."""
    folder = SCENARIOS_DIR / "02_engine_selection"
    folder.mkdir(parents=True, exist_ok=True)

    test_docx = folder / "engine_test.docx"
    doc = docx.Document()
    doc.add_heading("Проверка механизмов выбора движка", level=1)
    doc.add_paragraph("Тестирование параметров --engine auto, word, libreoffice.")
    doc.save(str(test_docx))

    all_ok = True
    notes = []

    # 1. Test auto engine
    auto_pdf = folder / "engine_auto.pdf"
    code_auto, stdout_auto, stderr_auto = run_command([
        str(CONVERT_SCRIPT),
        str(test_docx),
        "-o",
        str(auto_pdf),
        "--engine",
        "auto",
        "--force",
    ])
    valid_auto, reason_auto = is_valid_pdf_file(auto_pdf)
    if code_auto != 0 or not valid_auto:
        all_ok = False
        notes.append(f"Auto engine failed: code={code_auto}, {reason_auto}")

    # 2. Test invalid engine argument (argparse should reject with error code 2)
    inv_pdf = folder / "engine_invalid.pdf"
    code_inv, stdout_inv, stderr_inv = run_command([
        str(CONVERT_SCRIPT),
        str(test_docx),
        "-o",
        str(inv_pdf),
        "--engine",
        "nonexistent_engine",
    ])
    if code_inv == 0:
        all_ok = False
        notes.append("Invalid engine argument was unexpectedly accepted")

    # 3. Test explicit libreoffice engine
    lo_pdf = folder / "engine_lo.pdf"
    code_lo, stdout_lo, stderr_lo = run_command([
        str(CONVERT_SCRIPT),
        str(test_docx),
        "-o",
        str(lo_pdf),
        "--engine",
        "libreoffice",
        "--force",
    ])
    # LibreOffice might or might not be installed; if not, exit code must be 1 with descriptive message
    if code_lo == 0:
        valid_lo, _ = is_valid_pdf_file(lo_pdf)
        if not valid_lo:
            all_ok = False
            notes.append("LibreOffice returned code 0 but PDF is invalid")
    else:
        if "not found" not in stderr_lo.lower() and "error" not in stderr_lo.lower():
            all_ok = False
            notes.append(f"LibreOffice failed without clear error: {stderr_lo.strip()}")

    report = f"""# Сценарий 02: Выбор движков рендеринга

- Режим `--engine auto`: {'УСПЕХ' if code_auto == 0 and valid_auto else 'ОШИБКА'} (код: {code_auto})
- Режим `--engine invalid_engine`: {'УСПЕХ (отклонён валидатором)' if code_inv != 0 else 'ОШИБКА'}
- Режим `--engine libreoffice`: код возврата {code_lo} ({'сконвертирован' if code_lo == 0 else 'корректное сообщение об отсутствии/ошибке'})
- Замечания: {'; '.join(notes) if notes else 'Нет'}
- Статус: {'УСПЕХ' if all_ok else 'ОШИБКА'}
"""
    (folder / "scenario_report.md").write_text(report, encoding="utf-8")
    record_result("Scenario 02: Engine Selection", all_ok, "Verified auto mode, engine validation, and LibreOffice handling", folder)


# =========================================================================
# Scenario 03: Custom Output Path and Nested Destination Directory Creation
# =========================================================================
def scenario_03_custom_output_path():
    """Scenario 03: Test -o with deeply nested non-existent directory."""
    folder = SCENARIOS_DIR / "03_custom_output_path"
    folder.mkdir(parents=True, exist_ok=True)

    test_docx = folder / "source_document.docx"
    doc = docx.Document()
    doc.add_heading("Тест вложенных путей", level=1)
    doc.add_paragraph("Проверка корректного создания родительских каталогов при экспорте.")
    doc.save(str(test_docx))

    nested_target = folder / "sub1" / "sub2" / "deep_folder" / "nested_result.pdf"
    if nested_target.exists():
        nested_target.unlink()
    if nested_target.parent.exists():
        shutil.rmtree(str(folder / "sub1"), ignore_errors=True)

    code, stdout, stderr = run_command([
        str(CONVERT_SCRIPT),
        str(test_docx),
        "-o",
        str(nested_target),
        "--force",
    ])

    all_ok = True
    notes = []

    if code != 0:
        all_ok = False
        notes.append(f"CLI exited with code {code}, stderr: {stderr.strip()}")

    if not nested_target.is_file():
        all_ok = False
        notes.append(f"Nested output file not created: {nested_target}")
    else:
        valid, reason = is_valid_pdf_file(nested_target)
        if not valid:
            all_ok = False
            notes.append(f"Invalid PDF: {reason}")

    report = f"""# Сценарий 03: Создание вложенных путей назначения

- Исходный файл: `{test_docx.name}`
- Целевой путь: `{nested_target.relative_to(REPO_ROOT)}`
- Код возврата: {code}
- Создан ли каталог: {nested_target.parent.is_dir()}
- Валидность PDF: {'Да' if nested_target.is_file() else 'Нет'}
- Замечания: {'; '.join(notes) if notes else 'Нет'}
- Статус: {'УСПЕХ' if all_ok else 'ОШИБКА'}
"""
    (folder / "scenario_report.md").write_text(report, encoding="utf-8")
    record_result("Scenario 03: Custom Output & Nested Dirs", all_ok, "Verified nested path creation and target naming", folder)


# =========================================================================
# Scenario 04: Batch Directory Conversion
# =========================================================================
def scenario_04_batch_conversion():
    """Scenario 04: Batch convert an entire folder of DOCX files."""
    folder = SCENARIOS_DIR / "04_batch_conversion"
    batch_in = folder / "batch_in"
    batch_out = folder / "batch_out"

    if batch_in.exists():
        shutil.rmtree(str(batch_in), ignore_errors=True)
    if batch_out.exists():
        shutil.rmtree(str(batch_out), ignore_errors=True)

    batch_in.mkdir(parents=True, exist_ok=True)
    batch_out.mkdir(parents=True, exist_ok=True)

    doc_names = ["batch_doc_1.docx", "batch_doc_2.docx", "batch_doc_3.docx"]
    for i, name in enumerate(doc_names, 1):
        d = docx.Document()
        d.add_heading(f"Пакетный документ №{i}", level=1)
        d.add_paragraph(f"Содержимое документа для пакетного тестирования {i}.")
        d.save(str(batch_in / name))

    code, stdout, stderr = run_command([
        str(CONVERT_SCRIPT),
        str(batch_in),
        "-o",
        str(batch_out),
        "--force",
    ])

    all_ok = True
    notes = []

    if code != 0:
        all_ok = False
        notes.append(f"CLI exit code {code}, stderr: {stderr.strip()}")

    created_pdfs = list(batch_out.glob("*.pdf"))
    if len(created_pdfs) != len(doc_names):
        all_ok = False
        notes.append(f"Expected {len(doc_names)} PDFs, found {len(created_pdfs)}")

    for pdf in created_pdfs:
        valid, reason = is_valid_pdf_file(pdf)
        if not valid:
            all_ok = False
            notes.append(f"Batch PDF {pdf.name} invalid: {reason}")

    if "3 succeeded, 0 failed" not in stdout and "3 succeeded" not in stdout:
        notes.append("Expected batch summary in stdout not found")

    report = f"""# Сценарий 04: Пакетная конвертация каталога

- Входная папка: `{batch_in.name}` ({len(doc_names)} файлов)
- Выходная папка: `{batch_out.name}`
- Сгенерировано PDF: {len(created_pdfs)}
- Код возврата: {code}
- Замечания: {'; '.join(notes) if notes else 'Нет'}
- Статус: {'УСПЕХ' if all_ok else 'ОШИБКА'}
"""
    (folder / "scenario_report.md").write_text(report, encoding="utf-8")
    record_result("Scenario 04: Batch Conversion", all_ok, f"Converted {len(created_pdfs)} files in batch mode", folder)


# =========================================================================
# Scenario 05: Overwrite Protection and --force Flag
# =========================================================================
def scenario_05_overwrite_and_force():
    """Scenario 05: Verify overwrite prevention without --force and success with --force."""
    folder = SCENARIOS_DIR / "05_overwrite_and_force"
    folder.mkdir(parents=True, exist_ok=True)

    test_docx = folder / "overwrite_test.docx"
    doc = docx.Document()
    doc.add_heading("Тест перезаписи файлов", level=1)
    doc.add_paragraph("Проверка флага защиты от перезаписи --force.")
    doc.save(str(test_docx))

    target_pdf = folder / "overwrite_test.pdf"

    # Step 1: First conversion -> target_pdf exists
    code1, _, stderr1 = run_command([str(CONVERT_SCRIPT), str(test_docx), "-o", str(target_pdf), "--force"])
    if code1 != 0 or not target_pdf.is_file():
        record_result("Scenario 05: Overwrite & Force", False, "Initial PDF generation failed", folder)
        return

    orig_mtime = target_pdf.stat().st_mtime

    time.sleep(1.0)

    # Step 2: Second conversion WITHOUT --force -> Must fail
    code2, stdout2, stderr2 = run_command([str(CONVERT_SCRIPT), str(test_docx), "-o", str(target_pdf)])
    all_ok = True
    notes = []

    if code2 == 0:
        all_ok = False
        notes.append("Conversion succeeded without --force when output file already exists")
    if "already exists" not in stderr2.lower() and "already exists" not in stdout2.lower():
        notes.append("Error message did not mention file already exists")

    # Step 3: Third conversion WITH --force -> Must succeed and update file
    code3, stdout3, stderr3 = run_command([str(CONVERT_SCRIPT), str(test_docx), "-o", str(target_pdf), "--force"])
    if code3 != 0:
        all_ok = False
        notes.append(f"Conversion with --force failed: {stderr3.strip()}")
    new_mtime = target_pdf.stat().st_mtime
    if new_mtime <= orig_mtime:
        notes.append("File timestamp was not updated with --force")

    report = f"""# Сценарий 05: Защита от перезаписи и флаг --force

- Попытка без `--force` на существующий файл: {'Заблокирована (УСПЕХ)' if code2 != 0 else 'ОШИБКА'}
- Сообщение блокировки: `{stderr2.strip()}`
- Попытка с `--force`: {'УСПЕХ' if code3 == 0 else 'ОШИБКА'}
- Замечания: {'; '.join(notes) if notes else 'Нет'}
- Статус: {'УСПЕХ' if all_ok else 'ОШИБКА'}
"""
    (folder / "scenario_report.md").write_text(report, encoding="utf-8")
    record_result("Scenario 05: Overwrite Protection & Force", all_ok, "Verified collision protection and --force override", folder)


# =========================================================================
# Scenario 06: Filtering Word Lock Files (~$*.docx)
# =========================================================================
def scenario_06_lock_files_filtering():
    """Scenario 06: Verify exclusion of temporary lock files (~$*.docx)."""
    folder = SCENARIOS_DIR / "06_lock_files_filtering"
    in_dir = folder / "lock_test_dir"
    out_dir = folder / "lock_test_out"

    if in_dir.exists():
        shutil.rmtree(str(in_dir), ignore_errors=True)
    if out_dir.exists():
        shutil.rmtree(str(out_dir), ignore_errors=True)

    in_dir.mkdir(parents=True, exist_ok=True)
    out_dir.mkdir(parents=True, exist_ok=True)

    # 1. Create a valid file
    valid_doc = in_dir / "actual_document.docx"
    d = docx.Document()
    d.add_heading("Реальный документ", level=1)
    d.add_paragraph("Тестовый текст.")
    d.save(str(valid_doc))

    # 2. Create fake temporary lock file (~$actual_document.docx)
    lock_file = in_dir / "~$actual_document.docx"
    lock_file.write_bytes(b"Fake Word lock file header and author info")

    code, stdout, stderr = run_command([
        str(CONVERT_SCRIPT),
        str(in_dir),
        "-o",
        str(out_dir),
        "--force",
    ])

    all_ok = True
    notes = []

    if code != 0:
        all_ok = False
        notes.append(f"Exit code {code}, stderr: {stderr.strip()}")

    created_pdfs = list(out_dir.glob("*.pdf"))
    if len(created_pdfs) != 1:
        all_ok = False
        notes.append(f"Expected exactly 1 PDF, found {len(created_pdfs)}: {[p.name for p in created_pdfs]}")

    if (out_dir / "~$actual_document.pdf").exists():
        all_ok = False
        notes.append("Lock file was mistakenly processed into PDF")

    # Direct conversion of lock file must also be rejected
    code_direct, _, stderr_direct = run_command([
        str(CONVERT_SCRIPT),
        str(lock_file),
        "-o",
        str(out_dir / "direct_lock.pdf"),
    ])
    if code_direct == 0:
        all_ok = False
        notes.append("Direct conversion of ~$ lock file unexpectedly succeeded")

    report = f"""# Сценарий 06: Игнорирование временных lock-файлов Word

- Валидный документ: `{valid_doc.name}`
- Файл блокировки: `{lock_file.name}`
- Итоговых PDF в выходной папке: {len(created_pdfs)}
- Прямой вызов lock-файла: {'Отклонён (УСПЕХ)' if code_direct != 0 else 'ОШИБКА'}
- Замечания: {'; '.join(notes) if notes else 'Нет'}
- Статус: {'УСПЕХ' if all_ok else 'ОШИБКА'}
"""
    (folder / "scenario_report.md").write_text(report, encoding="utf-8")
    record_result("Scenario 06: Lock Files Filtering", all_ok, "Verified ~$*.docx lock files are ignored and rejected", folder)


# =========================================================================
# Scenario 07: Error Handling and Edge Cases (Corrupted, 0-byte, Missing Files)
# =========================================================================
def scenario_07_edge_cases_and_errors():
    """Scenario 07: Graceful error handling for corrupted, empty, or missing inputs."""
    folder = SCENARIOS_DIR / "07_edge_cases_and_errors"
    folder.mkdir(parents=True, exist_ok=True)

    all_ok = True
    notes = []

    # 1. 0-byte file
    zero_docx = folder / "zero_bytes.docx"
    zero_docx.write_bytes(b"")
    c_zero, out_zero, err_zero = run_command([str(CONVERT_SCRIPT), str(zero_docx)])
    if c_zero == 0:
        all_ok = False
        notes.append("0-byte file did not trigger error")

    # 2. Corrupted file (random text file renamed to .docx)
    corrupt_docx = folder / "corrupted.docx"
    corrupt_docx.write_text("This is not a valid zip archive or docx file.", encoding="utf-8")
    c_corrupt, out_corrupt, err_corrupt = run_command([str(CONVERT_SCRIPT), str(corrupt_docx)])
    if c_corrupt == 0:
        all_ok = False
        notes.append("Corrupted docx did not trigger error")

    # 3. Non-existent file
    missing_docx = folder / "does_not_exist_98765.docx"
    c_missing, out_missing, err_missing = run_command([str(CONVERT_SCRIPT), str(missing_docx)])
    if c_missing == 0:
        all_ok = False
        notes.append("Missing file did not trigger error")

    # Check for unhandled exceptions (tracebacks)
    combined_err = f"{err_zero}\n{err_corrupt}\n{err_missing}"
    if "Traceback (most recent call last):" in combined_err:
        all_ok = False
        notes.append("Unhandled python traceback detected in error responses")

    report = f"""# Сценарий 07: Обработка пограничных случаев и ошибок

- 0-байтовый файл: {'Отработан корректно (код 1)' if c_zero != 0 else 'ОШИБКА'}
- Повреждённый ZIP/DOCX: {'Отработан корректно (код 1)' if c_corrupt != 0 else 'ОШИБКА'}
- Несуществующий файл: {'Отработан корректно (код 1)' if c_missing != 0 else 'ОШИБКА'}
- Отсутствие необработанных Traceback: {'Да' if 'Traceback' not in combined_err else 'Нет'}
- Замечания: {'; '.join(notes) if notes else 'Нет'}
- Статус: {'УСПЕХ' if all_ok else 'ОШИБКА'}
"""
    (folder / "scenario_report.md").write_text(report, encoding="utf-8")
    record_result("Scenario 07: Edge Cases & Error Handling", all_ok, "Verified 0-byte, corrupted, and missing inputs", folder)


# =========================================================================
# Scenario 08: Real-World Samples Suite (file-sample_1, 2, 3)
# =========================================================================
def scenario_08_real_world_samples():
    """Scenario 08: End-to-end conversion of real-world sample documents."""
    folder = SCENARIOS_DIR / "08_real_world_samples"
    folder.mkdir(parents=True, exist_ok=True)

    samples = [SAMPLE_1, SAMPLE_2, SAMPLE_3]
    all_ok = True
    notes = []
    converted_count = 0

    for sample in samples:
        if not sample.is_file():
            notes.append(f"Sample not found: {sample.name}")
            continue

        out_pdf = folder / (sample.stem + ".pdf")
        t0 = time.time()
        code, stdout, stderr = run_command([
            str(CONVERT_SCRIPT),
            str(sample),
            "-o",
            str(out_pdf),
            "--force",
        ], timeout=120)
        elapsed = time.time() - t0

        if code != 0:
            all_ok = False
            notes.append(f"{sample.name}: Exit code {code}, stderr: {stderr.strip()}")
            continue

        valid, reason = is_valid_pdf_file(out_pdf)
        if not valid:
            all_ok = False
            notes.append(f"{sample.name}: Invalid PDF -> {reason}")
            continue

        try:
            pdf_doc = fitz.open(str(out_pdf))
            page_count = len(pdf_doc)
            pdf_doc.close()
            if page_count < 1:
                all_ok = False
                notes.append(f"{sample.name}: PDF has 0 pages")
            else:
                converted_count += 1
                print(f"  -> Converted {sample.name}: {page_count} pages in {elapsed:.1f}s ({out_pdf.stat().st_size} bytes)")
        except Exception as e:
            all_ok = False
            notes.append(f"{sample.name}: PyMuPDF inspection error: {e}")

    report = f"""# Сценарий 08: Тестирование на реальных документах (Samples)

- Проверено образцов: {len(samples)}
- Успешно сконвертировано: {converted_count}
- Замечания: {'; '.join(notes) if notes else 'Нет'}
- Статус: {'УСПЕХ' if all_ok and converted_count > 0 else 'ОШИБКА'}
"""
    (folder / "scenario_report.md").write_text(report, encoding="utf-8")
    record_result("Scenario 08: Real-World Samples Suite", all_ok, f"Converted {converted_count} real-world files successfully", folder)


# =========================================================================
# Scenario 09: PDF Quality & Content Inspection (PyMuPDF Text, Tables, Images)
# =========================================================================
def scenario_09_pdf_quality_validation():
    """Scenario 09: Deep inspection of PDF text encoding, Cyrillic glyphs, tables, and images."""
    folder = SCENARIOS_DIR / "09_pdf_quality_validation"
    folder.mkdir(parents=True, exist_ok=True)

    test_docx = folder / "quality_test.docx"
    doc = docx.Document()
    doc.add_heading("Анализ качества экспорта в PDF", level=0)

    # Specific Cyrillic characters to check for encoding regressions
    cyrillic_test_str = "Проверка символов: «ёлочка», спецзнак № 45/2026, тире — длинное, твердый знак: объём."
    doc.add_paragraph(cyrillic_test_str)

    # 4x3 table with numerical and textual data
    doc.add_heading("Спецификация модулей", level=1)
    table = doc.add_table(rows=4, cols=3)
    headers = ["Код", "Наименование модуля", "Статус проверки"]
    for i, h in enumerate(headers):
        table.cell(0, i).text = h
    rows = [
        ["MOD-01", "Ядро OpenXML", "Пройден"],
        ["MOD-02", "Движок рендеринга PDF", "Успешно"],
        ["MOD-03", "Валидатор PyMuPDF", "Активен"],
    ]
    for r_idx, row in enumerate(rows, 1):
        for c_idx, val in enumerate(row):
            table.cell(r_idx, c_idx).text = val

    # Embedded image
    img_path = folder / "chart_sample.png"
    create_sample_image(img_path, 450, 180, "График производительности")
    doc.add_heading("Встроенное изображение", level=1)
    doc.add_picture(str(img_path), width=Inches(4.5))

    # Page break to test multi-page documents
    doc.add_page_break()
    doc.add_heading("Вторая страница документа", level=1)
    doc.add_paragraph("Текст на второй странице для проверки разбивки страниц в PDF.")

    doc.save(str(test_docx))

    output_pdf = folder / "quality_test.pdf"
    code, stdout, stderr = run_command([
        str(CONVERT_SCRIPT),
        str(test_docx),
        "-o",
        str(output_pdf),
        "--force",
    ])

    all_ok = True
    notes = []
    full_text = ""
    images_found = 0
    pdf_pages = "N/A"

    if code != 0:
        all_ok = False
        notes.append(f"Exit code {code}, stderr: {stderr.strip()}")

    valid, reason = is_valid_pdf_file(output_pdf)
    if not valid:
        all_ok = False
        notes.append(f"PDF invalid: {reason}")
    else:
        try:
            pdf = fitz.open(str(output_pdf))
            pdf_pages = str(len(pdf))
            # 1. Verify page count >= 2
            if len(pdf) < 2:
                all_ok = False
                notes.append(f"Expected at least 2 pages, got {len(pdf)}")

            # 2. Verify Cyrillic text
            full_text = "".join(p.get_text() for p in pdf)
            if "объём" not in full_text or "«ёлочка»" not in full_text:
                all_ok = False
                notes.append("Cyrillic special characters missing or corrupted in PDF text layer")

            # 3. Verify Table content
            if "MOD-02" not in full_text or "Спецификация модулей" not in full_text:
                all_ok = False
                notes.append("Table content not found in PDF text layer")

            # 4. Verify Image rasterization
            images_found = 0
            for page in pdf:
                images_found += len(page.get_images())
            if images_found < 1:
                all_ok = False
                notes.append(f"Expected at least 1 embedded image in PDF, found {images_found}")

            pdf.close()
        except Exception as e:
            all_ok = False
            notes.append(f"PyMuPDF inspection error: {e}")

    report = f"""# Сценарий 09: Детальная валидация качества рендеринга PDF

- Исходный документ: `{test_docx.name}`
- Страниц в PDF: {pdf_pages} (ожидалось >= 2)
- Наличие изображений в PDF: {images_found}
- Сохранение кириллицы и спецсимволов: {'Да' if 'объём' in full_text else 'Нет'}
- Сохранение таблиц: {'Да' if 'MOD-02' in full_text else 'Нет'}
- Замечания: {'; '.join(notes) if notes else 'Нет'}
- Статус: {'УСПЕХ' if all_ok else 'ОШИБКА'}
"""
    (folder / "scenario_report.md").write_text(report, encoding="utf-8")
    record_result("Scenario 09: PDF Quality & Content Inspection", all_ok, "Verified multi-page layout, Cyrillic glyphs, table text, and embedded images", folder)


# =========================================================================
# Main Execution Runner
# =========================================================================
def main():
    print("=" * 80)
    print("ТЕСТОВЫЙ НАБОР ДЛЯ СКИЛА: docx-to-pdf")
    print(f"Целевой скрипт: {CONVERT_SCRIPT.relative_to(REPO_ROOT)}")
    print(f"Папка результатов: {SCENARIOS_DIR.relative_to(REPO_ROOT)}")
    print("=" * 80)

    scenario_01_single_file_word()
    scenario_02_engine_selection()
    scenario_03_custom_output_path()
    scenario_04_batch_conversion()
    scenario_05_overwrite_and_force()
    scenario_06_lock_files_filtering()
    scenario_07_edge_cases_and_errors()
    scenario_08_real_world_samples()
    scenario_09_pdf_quality_validation()

    # Generate test_summary.json
    all_passed = all(r["passed"] for r in test_results)
    summary = {
        "skill": "docx-to-pdf",
        "total_scenarios": len(test_results),
        "passed_scenarios": sum(1 for r in test_results if r["passed"]),
        "failed_scenarios": sum(1 for r in test_results if not r["passed"]),
        "all_passed": all_passed,
        "scenarios": test_results,
    }

    summary_file = TEST_ROOT / "test_summary.json"
    summary_file.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    print("\n" + "=" * 80)
    print(f"ИТОГИ ТЕСТИРОВАНИЯ docx-to-pdf: {summary['passed_scenarios']}/{summary['total_scenarios']} УСПЕШНО")
    print("=" * 80)

    for r in test_results:
        st = "✅ PASS" if r["passed"] else "❌ FAIL"
        print(f"[{st}] {r['scenario']}: {r['details']}")

    sys.exit(0 if all_passed else 1)


if __name__ == "__main__":
    main()
