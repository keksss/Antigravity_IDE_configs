# /// script
# requires-python = ">=3.10"
# dependencies = [
#     "python-docx>=1.1.2",
#     "mammoth>=1.8.0",
#     "markitdown[docx]>=0.0.1a4",
#     "Pillow>=10.0.0",
# ]
# ///
"""Convert Microsoft Word (.docx) documents to Markdown with full image, link, table, and comment extraction."""

import argparse
import mimetypes
import os
import re
import sys
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from docx import Document

# Ensure UTF-8 output on Windows consoles
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass


def sanitize_filename(name: str) -> str:
    """Sanitize string for safe filesystem usage."""
    return re.sub(r'[\\/*?:"<>| ]', "_", name)


def extract_comments_from_docx(docx_path: Path) -> list[dict]:
    """Extract comments from Word OpenXML with author, timestamp, comment text, and referenced fragment."""
    comments = []
    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    ns = {"w": w_ns}

    if not docx_path.is_file():
        return comments

    try:
        with zipfile.ZipFile(docx_path, "r") as z:
            namelist = z.namelist()
            if "word/comments.xml" not in namelist:
                return comments

            comments_xml = z.read("word/comments.xml")
            c_root = ET.fromstring(comments_xml)

            comments_map = {}
            for comment in c_root.findall(".//w:comment", ns):
                c_id = comment.get(f"{{{w_ns}}}id")
                author = comment.get(f"{{{w_ns}}}author") or "Unknown"
                date = comment.get(f"{{{w_ns}}}date") or ""
                initials = comment.get(f"{{{w_ns}}}initials") or ""

                p_texts = []
                for p in comment.findall(".//w:p", ns):
                    p_text = "".join([t.text for t in p.findall(".//w:t", ns) if t.text])
                    if p_text.strip():
                        p_texts.append(p_text.strip())
                text = "\n".join(p_texts) if p_texts else ""

                comments_map[c_id] = {
                    "id": c_id,
                    "author": author,
                    "date": date,
                    "initials": initials,
                    "text": text,
                    "referenced_text": "",
                }

            if "word/document.xml" in namelist:
                doc_xml = z.read("word/document.xml")
                d_root = ET.fromstring(doc_xml)

                active_comments = set()
                referenced_snippets = {cid: [] for cid in comments_map}
                paragraph_snippets = {cid: [] for cid in comments_map}

                for p_elem in d_root.findall(".//w:p", ns):
                    p_text = "".join([t.text for t in p_elem.findall(".//w:t", ns) if t.text])
                    for elem in p_elem.iter():
                        tag = elem.tag
                        if tag == f"{{{w_ns}}}commentRangeStart":
                            cid = elem.get(f"{{{w_ns}}}id")
                            if cid in comments_map:
                                active_comments.add(cid)
                        elif tag == f"{{{w_ns}}}commentRangeEnd":
                            cid = elem.get(f"{{{w_ns}}}id")
                            active_comments.discard(cid)
                        elif tag == f"{{{w_ns}}}commentReference":
                            cid = elem.get(f"{{{w_ns}}}id")
                            if cid in comments_map and p_text:
                                paragraph_snippets[cid].append(p_text.strip())
                        elif tag == f"{{{w_ns}}}t":
                            if elem.text:
                                for cid in active_comments:
                                    referenced_snippets[cid].append(elem.text)

                for cid in comments_map:
                    snip = "".join(referenced_snippets[cid]).strip()
                    if not snip and paragraph_snippets[cid]:
                        snip = paragraph_snippets[cid][0]
                    comments_map[cid]["referenced_text"] = snip

            comments = list(comments_map.values())
            try:
                comments.sort(key=lambda x: int(x["id"]))
            except ValueError:
                comments.sort(key=lambda x: str(x["id"]))

    except Exception as e:
        print(f"[Warning] Failed to extract comments: {e}", file=sys.stderr)

    return comments


def format_comments_markdown(comments: list[dict]) -> str:
    """Format extracted comments as a clean Markdown section."""
    if not comments:
        return ""
    lines = ["\n\n## Комментарии / Comments\n"]
    for c in comments:
        author = c.get("author") or "Аноним"
        date_str = f" *({c['date']})*" if c.get("date") else ""
        cid = c.get("id", "0")
        lines.append(f"- **[#{cid}] {author}**{date_str}:")
        if c.get("referenced_text"):
            lines.append(f"  > *К фрагменту:* \"{c['referenced_text']}\"\n  >")
        for text_line in c.get("text", "").splitlines():
            lines.append(f"  > {text_line}")
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def extract_media_from_docx(docx_path: Path, media_dir: Path, output_md_dir: Path = None) -> list[dict]:
    """Extract embedded images from Word document in document order and compute relative links."""
    media_dir.mkdir(parents=True, exist_ok=True)
    extracted = []
    if output_md_dir is None:
        output_md_dir = media_dir.parent

    try:
        doc = Document(str(docx_path))
        rels_map = {}
        for rel_id, rel in doc.part.rels.items():
            if "image" in rel.reltype:
                rels_map[rel_id] = rel

        if not rels_map:
            return extracted

        # Determine document order of images
        ordered_rids = []
        with zipfile.ZipFile(docx_path, "r") as z:
            if "word/document.xml" in z.namelist():
                root = ET.fromstring(z.read("word/document.xml"))
                a_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
                r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
                v_ns = "urn:schemas-microsoft-com:vml"

                for blip in root.findall(f".//{{{a_ns}}}blip"):
                    rid = blip.get(f"{{{r_ns}}}embed")
                    if rid and rid in rels_map and rid not in ordered_rids:
                        ordered_rids.append(rid)

                for img in root.findall(f".//{{{v_ns}}}imagedata"):
                    rid = img.get(f"{{{r_ns}}}id")
                    if rid and rid in rels_map and rid not in ordered_rids:
                        ordered_rids.append(rid)

        # Append any remaining image relationships
        for rid in rels_map:
            if rid not in ordered_rids:
                ordered_rids.append(rid)

        for idx, rid in enumerate(ordered_rids, start=1):
            rel = rels_map[rid]
            orig_name = Path(rel.target_ref).name
            ext = Path(orig_name).suffix or ".png"
            out_name = f"image_{idx:03d}{ext}"
            out_path = media_dir / out_name
            with open(out_path, "wb") as f:
                f.write(rel.target_part.blob)

            try:
                rel_path = os.path.relpath(out_path, output_md_dir).replace("\\", "/")
            except ValueError:
                rel_path = str(out_path).replace("\\", "/")

            extracted.append({
                "path": out_path,
                "rel_path": rel_path,
                "filename": out_name,
                "rid": rid
            })
    except Exception as e:
        print(f"[Warning] Failed to extract media: {e}", file=sys.stderr)

    return extracted


def convert_with_markitdown(docx_path: Path, output_md_path: Path, media_dir: Path) -> dict:
    """Convert DOCX to clean Markdown using Microsoft MarkItDown with relative image linking."""
    from markitdown import MarkItDown

    md = MarkItDown()
    result = md.convert(str(docx_path))
    md_text = result.text_content

    # Extract media files in document order
    extracted_images = extract_media_from_docx(docx_path, media_dir, output_md_path.parent)

    # Replace data:image/... URIs with relative image paths
    img_counter = [0]

    def image_replacer(match):
        alt_text = match.group(1)
        if img_counter[0] < len(extracted_images):
            item = extracted_images[img_counter[0]]
            img_counter[0] += 1
            alt = alt_text if alt_text else f"Image {img_counter[0]}"
            return f"![{alt}]({item['rel_path']})"
        return match.group(0)

    # Replace any ![alt](data:image/...) pattern
    pattern = r"!\[(.*?)\]\((data:image\/[^\)]+)\)"
    md_text = re.sub(pattern, image_replacer, md_text)

    output_md_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_md_path, "w", encoding="utf-8") as f:
        f.write(md_text.strip() + "\n")

    return {
        "engine": "markitdown",
        "docx_path": str(docx_path),
        "output_md": str(output_md_path),
        "media_dir": str(media_dir),
        "images_extracted": len(extracted_images),
        "warnings": [],
    }


def convert_with_mammoth(docx_path: Path, output_md_path: Path, media_dir: Path) -> dict:
    """Convert DOCX to Markdown using Mammoth with custom image handler and style mappings."""
    import mammoth

    media_dir.mkdir(parents=True, exist_ok=True)
    image_counter = 0
    extracted_images = []

    def image_handler(image):
        nonlocal image_counter
        image_counter += 1
        content_type = image.content_type
        extension = mimetypes.guess_extension(content_type) or ".png"
        if extension == ".jpe":
            extension = ".jpg"

        image_filename = f"image_{image_counter:03d}{extension}"
        image_file_path = media_dir / image_filename

        with image.open() as image_stream:
            with open(image_file_path, "wb") as out_img:
                out_img.write(image_stream.read())

        try:
            rel_image_path = os.path.relpath(image_file_path, output_md_path.parent)
        except ValueError:
            rel_image_path = str(image_file_path)

        rel_image_path = rel_image_path.replace("\\", "/")
        extracted_images.append(str(image_file_path))

        return {
            "src": rel_image_path,
            "alt": f"Image {image_counter}"
        }

    # Style mapping including Russian localized styles
    style_map = (
        "p[style-name='Heading 1'] => h1:fresh\n"
        "p[style-name='Heading 2'] => h2:fresh\n"
        "p[style-name='Heading 3'] => h3:fresh\n"
        "p[style-name='Heading 4'] => h4:fresh\n"
        "p[style-name='Heading 5'] => h5:fresh\n"
        "p[style-name='Heading 6'] => h6:fresh\n"
        "p[style-name='Заголовок 1'] => h1:fresh\n"
        "p[style-name='Заголовок 2'] => h2:fresh\n"
        "p[style-name='Заголовок 3'] => h3:fresh\n"
        "p[style-name='Заголовок 4'] => h4:fresh\n"
        "p[style-name='Заголовок 5'] => h5:fresh\n"
        "p[style-name='Заголовок 6'] => h6:fresh\n"
    )

    with open(docx_path, "rb") as docx_file:
        result = mammoth.convert_to_markdown(
            docx_file,
            convert_image=mammoth.images.inline(image_handler),
            style_map=style_map
        )
        md_text = result.value
        messages = result.messages

    md_text = re.sub(r'\n{3,}', '\n\n', md_text).strip() + "\n"

    output_md_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_md_path, "w", encoding="utf-8") as f:
        f.write(md_text)

    return {
        "engine": "mammoth",
        "docx_path": str(docx_path),
        "output_md": str(output_md_path),
        "media_dir": str(media_dir),
        "images_extracted": image_counter,
        "warnings": [str(m) for m in messages if m.type == "warning"],
    }


def extract_metadata(docx_path: Path) -> dict:
    """Extract core document properties, comments count, and style summary."""
    try:
        doc = Document(docx_path)
        core = doc.core_properties
        styles = sorted(list({p.style.name for p in doc.paragraphs if p.style}))
        comments = extract_comments_from_docx(docx_path)
        return {
            "title": core.title or "",
            "author": core.author or "",
            "created": str(core.created) if core.created else "",
            "modified": str(core.modified) if core.modified else "",
            "paragraphs_count": len(doc.paragraphs),
            "tables_count": len(doc.tables),
            "comments_count": len(comments),
            "styles_found": styles
        }
    except Exception as e:
        return {"error": f"Failed to extract metadata: {e}"}


def main():
    parser = argparse.ArgumentParser(
        description="Convert a .docx file to Markdown with full image, link, table, and comment extraction."
    )
    parser.add_argument("docx_file", type=Path, help="Path to input .docx document")
    parser.add_argument(
        "-o", "--output", type=Path, default=None,
        help="Path to output .md file (defaults to <docx_name>.md in same directory)"
    )
    parser.add_argument(
        "--media-dir", type=Path, default=None,
        help="Directory to save extracted images (defaults to <docx_name>_media)"
    )
    parser.add_argument(
        "--engine", choices=["auto", "markitdown", "mammoth"], default="auto",
        help="Conversion engine (default: auto; uses Microsoft MarkItDown with Mammoth fallback)"
    )
    parser.add_argument(
        "--include-headers", action="store_true",
        help="Include header and footer text in the output Markdown"
    )
    parser.add_argument(
        "--no-comments", action="store_true",
        help="Disable automatic extraction of Word comments into Markdown"
    )

    args = parser.parse_args()

    if not args.docx_file.is_file():
        print(f"Error: Input file not found: {args.docx_file}", file=sys.stderr)
        sys.exit(1)

    docx_path = args.docx_file.resolve()
    base_name = docx_path.stem

    output_md_path = args.output.resolve() if args.output else docx_path.parent / f"{base_name}.md"
    media_dir = args.media_dir.resolve() if args.media_dir else output_md_path.parent / f"{base_name}_media"

    res = None
    if args.engine in ("auto", "markitdown"):
        try:
            res = convert_with_markitdown(docx_path, output_md_path, media_dir)
        except Exception as e:
            if args.engine == "markitdown":
                print(f"Error: MarkItDown conversion failed: {e}", file=sys.stderr)
                sys.exit(1)
            print(f"[Warning] MarkItDown conversion failed ({e}), falling back to Mammoth...", file=sys.stderr)

    if res is None:
        res = convert_with_mammoth(docx_path, output_md_path, media_dir)

    # Read converted markdown
    with open(output_md_path, "r", encoding="utf-8") as f:
        md_content = f.read()

    # Optional: include header/footer text
    if args.include_headers:
        try:
            doc = Document(str(docx_path))
            header_texts = []
            for sec in doc.sections:
                for p in sec.header.paragraphs:
                    if p.text.strip():
                        header_texts.append(p.text.strip())
                for p in sec.footer.paragraphs:
                    if p.text.strip():
                        header_texts.append(p.text.strip())
            if header_texts:
                header_block = "> **Колонтитул:** " + " | ".join(header_texts) + "\n\n"
                md_content = header_block + md_content
        except Exception as e:
            print(f"[Warning] Failed to process headers: {e}", file=sys.stderr)

    # Optional: extract comments (enabled by default)
    comments_extracted_count = 0
    if not args.no_comments:
        comments = extract_comments_from_docx(docx_path)
        if comments:
            comments_extracted_count = len(comments)
            comments_block = format_comments_markdown(comments)
            md_content = md_content.rstrip() + comments_block

    # Write updated markdown
    with open(output_md_path, "w", encoding="utf-8") as f:
        f.write(md_content.strip() + "\n")

    metadata = extract_metadata(docx_path)

    print(f"SUCCESS: Converted '{docx_path.name}' to Markdown (engine: {res['engine']}).")
    print(f"  Output Markdown: {output_md_path}")
    print(f"  Media Folder:    {media_dir} ({res['images_extracted']} images extracted)")
    if comments_extracted_count > 0:
        print(f"  Comments:        {comments_extracted_count} comments extracted")
    if metadata.get("styles_found"):
        print(f"  Styles detected: {', '.join(metadata['styles_found'][:6])}...")
    if res.get("warnings"):
        print(f"  Warnings ({len(res['warnings'])}):")
        for w in res["warnings"][:3]:
            print(f"    - {w}")


if __name__ == "__main__":
    main()
