# /// script
# requires-python = ">=3.10"
# dependencies = [
#     "python-docx>=1.1.2",
# ]
# ///
"""Inspect structure, styles, tables, and content of a Word (.docx) document."""

import argparse
import json
import sys
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

# Ensure UTF-8 output on Windows consoles
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass


def get_table_column_count(table) -> int:
    """Safely calculate table column count without failing on merged cells."""
    # 1. Try table.columns, but catch NotImplementedError / IndexError
    try:
        return len(table.columns)
    except Exception:
        pass

    # 2. Check XML tblGrid definition
    try:
        tblGrid = table._tbl.find(qn("w:tblGrid"))
        if tblGrid is not None:
            gridCols = tblGrid.findall(qn("w:gridCol"))
            if gridCols:
                return len(gridCols)
    except Exception:
        pass

    # 3. Fallback to maximum number of cells across rows
    if table.rows:
        return max((len(r.cells) for r in table.rows), default=0)
    return 0


def get_heading_info(p, idx: int) -> dict | None:
    """Detect heading level across languages and custom styles."""
    text = p.text.strip()
    if not text:
        return None

    style_name = p.style.name if p.style else ""
    outline_lvl = None

    # Check XML outline level (w:outlineLvl) in paragraph properties
    if p._p.pPr is not None:
        ol = p._p.pPr.find(qn("w:outlineLvl"))
        if ol is not None and ol.get(qn("w:val")) is not None:
            try:
                outline_lvl = int(ol.get(qn("w:val")))
            except ValueError:
                pass

    # Check multilingual style names
    s_lower = style_name.lower()
    multilingual_keywords = [
        "heading", "заголовок", "titre", "überschrift", "rubrik",
        "encabezado", "header", "subheading", "подзаголовок"
    ]
    is_heading_style = any(kw in s_lower for kw in multilingual_keywords)

    # All-caps short line heuristic
    is_caps = text.isupper() and 2 < len(text) < 80

    if outline_lvl is not None or is_heading_style or is_caps:
        return {
            "index": idx,
            "style": style_name or "Custom",
            "outline_level": outline_lvl,
            "text": text
        }
    return None


def inspect_document(docx_path: Path, mode: str = "summary") -> dict:
    """Extract structural information from a docx file."""
    try:
        doc = Document(str(docx_path))
    except PermissionError:
        print(f"Error: Permission denied accessing '{docx_path.name}'. The file may be currently open in Microsoft Word. Please close Word and try again.", file=sys.stderr)
        sys.exit(1)
    except Exception as e:
        print(f"Error: Failed to open '{docx_path.name}': {e}", file=sys.stderr)
        sys.exit(1)

    # Core Properties
    core = doc.core_properties
    meta = {
        "title": core.title or "",
        "author": core.author or "",
        "created": str(core.created) if core.created else "",
        "modified": str(core.modified) if core.modified else "",
        "paragraphs_total": len(doc.paragraphs),
        "tables_total": len(doc.tables),
    }

    # Headings / Outline
    headings = []
    for idx, p in enumerate(doc.paragraphs):
        h_info = get_heading_info(p, idx)
        if h_info:
            headings.append(h_info)

    # Tables info (safe column counting)
    tables_info = []
    for idx, t in enumerate(doc.tables):
        header_row = []
        if t.rows:
            try:
                header_row = [cell.text.strip().replace("\n", " ") for cell in t.rows[0].cells]
            except Exception:
                pass

        cols_count = get_table_column_count(t)
        tables_info.append({
            "table_index": idx,
            "rows": len(t.rows),
            "columns": cols_count,
            "headers": header_row[:6]
        })

    # Header and Footer inspection
    sections_info = []
    for s_idx, sec in enumerate(doc.sections):
        hdr = "\n".join(p.text.strip() for p in sec.header.paragraphs if p.text.strip())
        ftr = "\n".join(p.text.strip() for p in sec.footer.paragraphs if p.text.strip())
        if hdr or ftr:
            sections_info.append({
                "section": s_idx,
                "header": hdr,
                "footer": ftr
            })

    # Images / media count via relationships
    images_count = sum(1 for rel in doc.part.rels.values() if "image" in rel.reltype)
    meta["images_total"] = images_count

    # All styles in use
    used_styles = sorted(list({p.style.name for p in doc.paragraphs if p.style}))

    if mode == "outline":
        return {"file": docx_path.name, "metadata": meta, "headings": headings}
    elif mode == "tables":
        return {"file": docx_path.name, "tables": tables_info}
    elif mode == "styles":
        return {"file": docx_path.name, "used_styles": used_styles}
    elif mode == "headers":
        return {"file": docx_path.name, "sections": sections_info}
    elif mode == "raw":
        all_text = []
        for p in doc.paragraphs:
            if p.text.strip():
                all_text.append(p.text)
        for t in doc.tables:
            for r in t.rows:
                row_str = " | ".join(c.text.strip() for c in r.cells if c.text.strip())
                if row_str:
                    all_text.append(row_str)
        return {"file": docx_path.name, "text": "\n".join(all_text)}
    else:  # summary
        return {
            "file": docx_path.name,
            "metadata": meta,
            "used_styles": used_styles[:10],
            "headings_count": len(headings),
            "headings_preview": headings[:8],
            "tables": tables_info,
            "sections_with_headers": len(sections_info),
        }


def main():
    parser = argparse.ArgumentParser(
        description="Inspect structure, headings, tables, and styles of a .docx file."
    )
    parser.add_argument("docx_file", type=Path, help="Path to .docx file")
    parser.add_argument(
        "--mode", choices=["summary", "outline", "styles", "tables", "headers", "raw"], default="summary",
        help="Inspection mode (default: summary)"
    )
    parser.add_argument("--json", action="store_true", help="Output results as JSON")

    args = parser.parse_args()

    if not args.docx_file.is_file():
        print(f"Error: File not found: {args.docx_file}", file=sys.stderr)
        sys.exit(1)

    result = inspect_document(args.docx_file.resolve(), mode=args.mode)

    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return

    # Formatted terminal output
    print(f"=== DOCX Inspection: {result['file']} ===")
    if "metadata" in result:
        m = result["metadata"]
        print(f"Author: {m['author'] or 'N/A'} | Modified: {m['modified'] or 'N/A'}")
        print(f"Total Paragraphs: {m['paragraphs_total']} | Total Tables: {m['tables_total']} | Images: {m.get('images_total', 0)}")

    if "used_styles" in result:
        print("\nStyles in Use:")
        for s in result["used_styles"]:
            print(f"  - {s}")

    if "headings" in result or "headings_preview" in result:
        h_list = result.get("headings") or result.get("headings_preview")
        print(f"\nDocument Headings ({len(h_list)}):")
        for h in h_list:
            lvl_str = f" [lvl {h['outline_level']}]" if h.get("outline_level") is not None else ""
            print(f"  [{h['style']}]{lvl_str} (p#{h['index']}) {h['text']}")

    if "tables" in result and result["tables"]:
        print(f"\nTables ({len(result['tables'])}):")
        for t in result["tables"]:
            hdr = " | ".join(t['headers']) if t['headers'] else 'No header'
            print(f"  Table #{t['table_index']}: {t['rows']} rows x {t['columns']} cols [{hdr}]")

    if "sections" in result:
        print(f"\nHeaders & Footers ({len(result['sections'])}):")
        for s in result["sections"]:
            if s["header"]:
                print(f"  Section #{s['section']} Header: {s['header']}")
            if s["footer"]:
                print(f"  Section #{s['section']} Footer: {s['footer']}")

    if "text" in result:
        print("\nRaw Text (Body & Tables):")
        print(result["text"])


if __name__ == "__main__":
    main()
