# /// script
# requires-python = ">=3.10"
# dependencies = [
#     "openpyxl>=3.1.2",
#     "pandas>=2.0.0",
# ]
# ///
"""Self-contained Excel & Tabular Data Inspection, Markdown, JSON, and HTML Converter.

Quickly inspects structure, sheets, dimensions, formulas, and converts tabular data from
.xlsx, .xlsm, .csv, and .tsv files into LLM-friendly formats (Markdown, JSON, JSON-Grid, HTML)
without requiring external heavy tools or cloud services.
"""

import argparse
import csv
import datetime
import decimal
import html
import json
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

# Ensure UTF-8 output on Windows consoles
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass


def json_serializer(obj: Any) -> Any:
    """Safe serializer for datetime, dates, decimals, and custom objects."""
    if isinstance(obj, (datetime.datetime, datetime.date, datetime.time)):
        return obj.isoformat()
    if isinstance(obj, decimal.Decimal):
        return float(obj)
    if hasattr(obj, "__str__"):
        return str(obj)
    return repr(obj)


def format_markdown_table(headers: List[str], rows: List[List[str]]) -> str:
    """Render a clean GitHub-flavored markdown table."""
    if not headers and not rows:
        return "*Empty table*"

    col_count = max(len(headers), max((len(r) for r in rows), default=0))
    padded_headers = headers + [""] * (col_count - len(headers))
    clean_headers = [h.replace("\n", " ").replace("|", "\\|") for h in padded_headers]

    lines = []
    lines.append("| " + " | ".join(clean_headers) + " |")
    lines.append("| " + " | ".join(["---"] * col_count) + " |")

    for r in rows:
        padded_row = r + [""] * (col_count - len(r))
        clean_row = [str(c).replace("\n", " ").replace("|", "\\|") if c is not None else "" for c in padded_row]
        lines.append("| " + " | ".join(clean_row) + " |")

    return "\n".join(lines)


def get_merged_cells_map(ws, row_limit: int, col_limit: int) -> Tuple[Dict[Tuple[int, int], Tuple[int, int]], Set[Tuple[int, int]], List[str]]:
    """Map merged cell ranges to top-left anchors (rowspan, colspan) and non-anchor covered cells."""
    anchors: Dict[Tuple[int, int], Tuple[int, int]] = {}
    covered: Set[Tuple[int, int]] = set()
    merged_str_list: List[str] = []

    if not hasattr(ws, "merged_cells") or not hasattr(ws.merged_cells, "ranges"):
        return anchors, covered, merged_str_list

    for rng in ws.merged_cells.ranges:
        merged_str_list.append(str(rng))
        min_r, min_c = rng.min_row, rng.min_col
        max_r, max_c = rng.max_row, rng.max_col

        if min_r > row_limit or min_c > col_limit:
            continue

        eff_max_r = min(max_r, row_limit)
        eff_max_c = min(max_c, col_limit)
        rowspan = eff_max_r - min_r + 1
        colspan = eff_max_c - min_c + 1

        anchors[(min_r, min_c)] = (rowspan, colspan)
        for r in range(min_r, eff_max_r + 1):
            for c in range(min_c, eff_max_c + 1):
                if (r, c) != (min_r, min_c):
                    covered.add((r, c))

    return anchors, covered, merged_str_list


def convert_sheet_to_json_grid(ws_form, ws_val, sheet_name: str, max_rows: int, max_cols: int, dense: bool) -> Dict[str, Any]:
    """Convert worksheet into a cell-by-cell coordinate map with values and formulas."""
    from openpyxl.utils import get_column_letter

    row_limit = ws_form.max_row or 0
    if max_rows > 0:
        row_limit = min(row_limit, max_rows)

    col_limit = ws_form.max_column or 0
    if max_cols > 0:
        col_limit = min(col_limit, max_cols)

    dim_str = f"A1:{get_column_letter(col_limit)}{row_limit}" if row_limit > 0 and col_limit > 0 else "Empty"
    _, _, merged_ranges = get_merged_cells_map(ws_form, row_limit, col_limit)

    cells: Dict[str, Dict[str, Any]] = {}
    for r in range(1, row_limit + 1):
        for c in range(1, col_limit + 1):
            coord = f"{get_column_letter(c)}{r}"
            cell_form = ws_form.cell(row=r, column=c).value
            cell_val = ws_val.cell(row=r, column=c).value

            if dense and cell_form is None and cell_val is None:
                continue

            entry: Dict[str, Any] = {}
            if cell_val is not None:
                entry["v"] = cell_val
            if isinstance(cell_form, str) and cell_form.startswith("="):
                entry["f"] = cell_form
            if entry:
                cells[coord] = entry

    return {
        "sheet": sheet_name,
        "dimensions": dim_str,
        "total_rows": ws_form.max_row or 0,
        "total_columns": ws_form.max_column or 0,
        "exported_rows": row_limit,
        "exported_columns": col_limit,
        "merged_ranges": merged_ranges,
        "cells": cells,
    }


def convert_sheet_to_json_records(ws_form, ws_val, sheet_name: str, max_rows: int, max_cols: int, dense: bool) -> Dict[str, Any]:
    """Convert worksheet into tabular JSON records with header detection."""
    from openpyxl.utils import get_column_letter

    row_limit = ws_form.max_row or 0
    if max_rows > 0:
        row_limit = min(row_limit, max_rows)

    col_limit = ws_form.max_column or 0
    if max_cols > 0:
        col_limit = min(col_limit, max_cols)

    if row_limit == 0 or col_limit == 0:
        return {"sheet": sheet_name, "records": []}

    # Extract raw rows
    raw_rows: List[List[Any]] = []
    formula_map: Dict[str, str] = {}

    for r in range(1, row_limit + 1):
        row_vals = []
        for c in range(1, col_limit + 1):
            val = ws_val.cell(row=r, column=c).value
            form = ws_form.cell(row=r, column=c).value
            row_vals.append(val)
            if isinstance(form, str) and form.startswith("="):
                coord = f"{get_column_letter(c)}{r}"
                formula_map[coord] = form
        raw_rows.append(row_vals)

    # Check if first row contains non-empty strings suitable as column headers
    first_row = raw_rows[0]
    has_headers = any(first_row) and all(isinstance(x, str) and x.strip() for x in first_row if x is not None)

    records: List[Dict[str, Any]] = []
    if has_headers and len(raw_rows) > 1:
        headers = [str(h).strip() if h is not None else f"Col_{get_column_letter(i+1)}" for i, h in enumerate(first_row)]
        # Ensure unique headers
        seen_headers: Dict[str, int] = {}
        unique_headers = []
        for h in headers:
            if h in seen_headers:
                seen_headers[h] += 1
                unique_headers.append(f"{h}_{seen_headers[h]}")
            else:
                seen_headers[h] = 0
                unique_headers.append(h)

        for r_idx, row in enumerate(raw_rows[1:], start=2):
            record = {}
            for col_idx, val in enumerate(row):
                if dense and (val is None or val == ""):
                    continue
                header_name = unique_headers[col_idx] if col_idx < len(unique_headers) else f"Col_{col_idx+1}"
                record[header_name] = val
            if record:
                records.append(record)
    else:
        # Fallback to indexed row records with column letters
        for r_idx, row in enumerate(raw_rows, start=1):
            record = {"_row": r_idx}
            for col_idx, val in enumerate(row):
                if dense and (val is None or val == ""):
                    continue
                col_letter = get_column_letter(col_idx + 1)
                record[col_letter] = val
            records.append(record)

    return {
        "sheet": sheet_name,
        "total_rows": ws_form.max_row or 0,
        "total_columns": ws_form.max_column or 0,
        "exported_rows": len(records),
        "records": records,
        "formulas": formula_map if formula_map else None,
    }


def convert_sheet_to_html_table(ws_form, ws_val, sheet_name: str, max_rows: int, max_cols: int, show_formulas: bool) -> str:
    """Render a semantic HTML <table> with merged cell handling (colspan/rowspan)."""
    from openpyxl.utils import get_column_letter

    row_limit = ws_form.max_row or 0
    if max_rows > 0:
        row_limit = min(row_limit, max_rows)

    col_limit = ws_form.max_column or 0
    if max_cols > 0:
        col_limit = min(col_limit, max_cols)

    if row_limit == 0 or col_limit == 0:
        return f'<table class="excel-table"><caption>Sheet: {html.escape(sheet_name)}</caption><tr><td><em>Empty sheet</em></td></tr></table>'

    anchors, covered, _ = get_merged_cells_map(ws_form, row_limit, col_limit)

    lines = []
    lines.append(f'<table class="excel-table" data-sheet="{html.escape(sheet_name)}">')
    lines.append(f'  <caption>Sheet: <strong>{html.escape(sheet_name)}</strong> (Rows 1–{row_limit}, Cols 1–{col_limit})</caption>')
    lines.append('  <thead>')
    lines.append('    <tr>')
    lines.append('      <th class="row-header">#</th>')
    for c in range(1, col_limit + 1):
        col_letter = get_column_letter(c)
        lines.append(f'      <th class="col-header">{col_letter}</th>')
    lines.append('    </tr>')
    lines.append('  </thead>')
    lines.append('  <tbody>')

    for r in range(1, row_limit + 1):
        lines.append('    <tr>')
        lines.append(f'      <td class="row-header">{r}</td>')
        for c in range(1, col_limit + 1):
            if (r, c) in covered:
                continue

            cell_form = ws_form.cell(row=r, column=c).value
            cell_val = ws_val.cell(row=r, column=c).value

            # Determine cell attributes
            attrs = []
            if (r, c) in anchors:
                rowspan, colspan = anchors[(r, c)]
                if rowspan > 1:
                    attrs.append(f'rowspan="{rowspan}"')
                if colspan > 1:
                    attrs.append(f'colspan="{colspan}"')

            # Class for data type alignment
            is_num = isinstance(cell_val, (int, float, decimal.Decimal))
            css_class = "num" if is_num else "text"
            attrs.append(f'class="{css_class}"')

            # Cell display text
            val_display = ""
            if cell_val is not None:
                val_display = html.escape(str(cell_val))

            formula_span = ""
            if isinstance(cell_form, str) and cell_form.startswith("="):
                form_esc = html.escape(cell_form)
                attrs.append(f'data-formula="{form_esc}"')
                if show_formulas:
                    val_display = f'<code>{form_esc}</code>'
                else:
                    formula_span = f'<span class="formula-tag" title="Formula: {form_esc}">fx</span>'

            inner = val_display + formula_span
            attr_str = (" " + " ".join(attrs)) if attrs else ""
            lines.append(f'      <td{attr_str}>{inner}</td>')
        lines.append('    </tr>')

    lines.append('  </tbody>')
    lines.append('</table>')
    return "\n".join(lines)


HTML_PAGE_TEMPLATE = """<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>{title}</title>
  <style>
    :root {{
      --font-sans: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
      --bg: #f8fafc;
      --card-bg: #ffffff;
      --text: #1e293b;
      --border: #cbd5e1;
      --header-bg: #f1f5f9;
      --accent: #2563eb;
    }}
    body {{
      font-family: var(--font-sans);
      background-color: var(--bg);
      color: var(--text);
      margin: 0;
      padding: 24px;
      line-height: 1.5;
    }}
    header {{
      margin-bottom: 24px;
      padding-bottom: 12px;
      border-bottom: 2px solid var(--border);
    }}
    h1 {{
      margin: 0 0 8px 0;
      font-size: 1.5rem;
    }}
    .meta {{
      font-size: 0.875rem;
      color: #64748b;
    }}
    .sheet-section {{
      margin-bottom: 32px;
    }}
    .excel-table {{
      border-collapse: collapse;
      width: 100%;
      background: var(--card-bg);
      border: 1px solid var(--border);
      border-radius: 6px;
      overflow: hidden;
      box-shadow: 0 1px 3px rgba(0,0,0,0.05);
      margin-bottom: 16px;
    }}
    .excel-table caption {{
      font-size: 1rem;
      font-weight: 600;
      text-align: left;
      padding: 10px 14px;
      background: var(--header-bg);
      border-bottom: 1px solid var(--border);
    }}
    .excel-table th, .excel-table td {{
      border: 1px solid var(--border);
      padding: 6px 12px;
      font-size: 0.875rem;
    }}
    .excel-table th.col-header, .excel-table td.row-header {{
      background: #f8fafc;
      color: #64748b;
      font-weight: 600;
      text-align: center;
      user-select: none;
      width: 40px;
    }}
    .excel-table td.num {{
      text-align: right;
      font-variant-numeric: tabular-nums;
    }}
    .excel-table td.text {{
      text-align: left;
    }}
    .excel-table tr:hover td:not(.row-header) {{
      background-color: #f1f5f9;
    }}
    .formula-tag {{
      display: inline-block;
      font-size: 0.65rem;
      font-weight: 700;
      color: #0284c7;
      background: #e0f2fe;
      border: 1px solid #bae6fd;
      border-radius: 3px;
      padding: 0 3px;
      margin-left: 6px;
      vertical-align: middle;
      cursor: help;
    }}
    code {{
      font-family: ui-monospace, SFMono-Regular, Menlo, Monaco, Consolas, monospace;
      font-size: 0.8rem;
      color: #0f766e;
    }}
  </style>
</head>
<body>
  <header>
    <h1>Workbook: {file_name}</h1>
    <div class="meta">Exported for LLM & browser inspection &bull; Generated on {date}</div>
  </header>
  <main>
    {content}
  </main>
</body>
</html>
"""


def output_result(content: str, out_path: Optional[Path]) -> None:
    """Print to stdout or write directly to target file."""
    if out_path:
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_text(content, encoding="utf-8")
        print(f"Successfully exported to {out_path} ({len(content.encode('utf-8'))} bytes)")
    else:
        sys.stdout.write(content)
        if not content.endswith("\n"):
            sys.stdout.write("\n")


def inspect_csv(file_path: Path, mode: str, fmt: str, max_rows: int, out_path: Optional[Path]) -> None:
    """Inspect CSV or TSV file and export in requested format."""
    delimiter = "\t" if file_path.suffix.lower() == ".tsv" else ","
    rows = []
    try:
        with file_path.open("r", encoding="utf-8-sig", errors="replace") as f:
            reader = csv.reader(f, delimiter=delimiter)
            for i, row in enumerate(reader):
                if max_rows > 0 and i >= max_rows + 1:
                    break
                rows.append(row)
    except Exception as e:
        print(f"Error reading {file_path}: {e}", file=sys.stderr)
        sys.exit(1)

    if not rows:
        output_result(f"File {file_path.name} is empty.", out_path)
        return

    if fmt == "json":
        headers = rows[0]
        records = []
        for r in rows[1:]:
            record = {headers[i] if i < len(headers) else f"Col_{i+1}": val for i, val in enumerate(r)}
            records.append(record)
        res_json = json.dumps({"file": file_path.name, "records": records}, indent=2, default=json_serializer, ensure_ascii=False)
        output_result(res_json, out_path)

    elif fmt == "html":
        lines = [f'<table class="excel-table"><caption>File: {html.escape(file_path.name)}</caption><thead><tr>']
        for col_idx, h in enumerate(rows[0]):
            lines.append(f'<th class="col-header">{html.escape(h)}</th>')
        lines.append('</tr></thead><tbody>')
        for r in rows[1:]:
            lines.append('<tr>')
            for cell in r:
                lines.append(f'<td class="text">{html.escape(cell)}</td>')
            lines.append('</tr>')
        lines.append('</tbody></table>')
        full_html = HTML_PAGE_TEMPLATE.format(
            title=f"CSV Export: {file_path.name}",
            file_name=file_path.name,
            date=datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            content="\n".join(lines),
        )
        output_result(full_html, out_path)

    else:  # markdown
        headers = ["Row"] + [f"Col {idx+1}" for idx in range(len(rows[0]))]
        row_limit = min(len(rows), max_rows) if max_rows > 0 else len(rows)
        formatted_rows = [[f"{i+1}"] + r for i, r in enumerate(rows[:row_limit])]
        res_md = f"### {file_path.name} ({len(rows)} rows sampled)\n\n" + format_markdown_table(headers, formatted_rows)
        output_result(res_md, out_path)


def inspect_xlsx_summary(file_path: Path, fmt: str, out_path: Optional[Path]) -> None:
    """Provide a high-level summary of workbook sheets and contents."""
    from openpyxl import load_workbook

    wb = load_workbook(file_path, data_only=False, read_only=False)
    try:
        sheet_infos = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            formula_count = 0
            if hasattr(ws, "iter_rows"):
                for row in ws.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and cell.value.startswith("="):
                            formula_count += 1

            sheet_infos.append({
                "name": sheet_name,
                "dimensions": ws.dimensions or "Unknown",
                "max_row": ws.max_row or 0,
                "max_column": ws.max_column or 0,
                "formula_count": formula_count,
                "state": ws.sheet_state,
            })

        if fmt == "json":
            payload = {
                "file": file_path.name,
                "total_sheets": len(sheet_infos),
                "sheets": sheet_infos,
            }
            output_result(json.dumps(payload, indent=2, default=json_serializer, ensure_ascii=False), out_path)

        elif fmt == "html":
            table_lines = [
                f'<table class="excel-table"><caption>Workbook Summary: {html.escape(file_path.name)}</caption>',
                '<thead><tr><th>Sheet Name</th><th>Dimensions</th><th>Max Row</th><th>Max Col</th><th>Formula Count</th><th>State</th></tr></thead>',
                '<tbody>'
            ]
            for s in sheet_infos:
                table_lines.append(
                    f'<tr><td><strong>{html.escape(s["name"])}</strong></td>'
                    f'<td><code>{html.escape(s["dimensions"])}</code></td>'
                    f'<td class="num">{s["max_row"]}</td>'
                    f'<td class="num">{s["max_column"]}</td>'
                    f'<td class="num">{s["formula_count"]}</td>'
                    f'<td>{html.escape(s["state"])}</td></tr>'
                )
            table_lines.append('</tbody></table>')
            full_html = HTML_PAGE_TEMPLATE.format(
                title=f"Summary: {file_path.name}",
                file_name=file_path.name,
                date=datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                content="\n".join(table_lines),
            )
            output_result(full_html, out_path)

        else:  # markdown
            summary_headers = ["Sheet Name", "Dimensions", "Max Row", "Max Col", "Formula Count", "State"]
            summary_rows = [
                [s["name"], s["dimensions"], str(s["max_row"]), str(s["max_column"]), str(s["formula_count"]), s["state"]]
                for s in sheet_infos
            ]
            md_out = (
                f"## Workbook Summary: `{file_path.name}`\n\n"
                f"- **Total Sheets:** {len(wb.sheetnames)}\n"
                f"- **Sheets:** {', '.join(wb.sheetnames)}\n\n"
                + format_markdown_table(summary_headers, summary_rows)
            )
            output_result(md_out, out_path)

    finally:
        wb.close()


def inspect_xlsx_preview(
    file_path: Path,
    target_sheet: Optional[str],
    all_sheets: bool,
    max_rows: int,
    max_cols: int,
    show_formulas: bool,
    fmt: str,
    dense: bool,
    out_path: Optional[Path],
) -> None:
    """Render tabular preview or export in Markdown, JSON, JSON-Grid, or HTML."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    wb_form = load_workbook(file_path, data_only=False)
    wb_val = load_workbook(file_path, data_only=True)

    try:
        sheets = wb_form.sheetnames if all_sheets else [target_sheet or wb_form.active.title]

        for s in sheets:
            if s not in wb_form.sheetnames:
                print(f"Error: Sheet '{s}' not found. Available: {', '.join(wb_form.sheetnames)}", file=sys.stderr)
                sys.exit(1)

        if fmt == "json":
            results = []
            for s in sheets:
                res = convert_sheet_to_json_records(wb_form[s], wb_val[s], s, max_rows, max_cols, dense)
                results.append(res)
            payload = results[0] if len(results) == 1 else {"file": file_path.name, "sheets": results}
            output_result(json.dumps(payload, indent=2, default=json_serializer, ensure_ascii=False), out_path)

        elif fmt == "json-grid":
            results = []
            for s in sheets:
                res = convert_sheet_to_json_grid(wb_form[s], wb_val[s], s, max_rows, max_cols, dense)
                results.append(res)
            payload = results[0] if len(results) == 1 else {"file": file_path.name, "sheets": results}
            output_result(json.dumps(payload, indent=2, default=json_serializer, ensure_ascii=False), out_path)

        elif fmt == "html":
            tables = []
            for s in sheets:
                t = convert_sheet_to_html_table(wb_form[s], wb_val[s], s, max_rows, max_cols, show_formulas)
                tables.append(f'<div class="sheet-section">\n{t}\n</div>')
            full_html = HTML_PAGE_TEMPLATE.format(
                title=f"Export: {file_path.name}",
                file_name=file_path.name,
                date=datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                content="\n".join(tables),
            )
            output_result(full_html, out_path)

        else:  # markdown
            md_sections = []
            for s in sheets:
                ws_form = wb_form[s]
                ws_val = wb_val[s]

                row_limit = ws_form.max_row or 0
                if max_rows > 0:
                    row_limit = min(row_limit, max_rows)

                col_limit = ws_form.max_column or 0
                if max_cols > 0:
                    col_limit = min(col_limit, max_cols)

                sec = [f"### Sheet: `{s}` (Rows 1–{row_limit}, Cols 1–{col_limit})\n"]
                if row_limit == 0 or col_limit == 0:
                    sec.append("*Sheet has no data.*")
                else:
                    headers = ["Row"] + [get_column_letter(c) for c in range(1, col_limit + 1)]
                    grid = []
                    for r in range(1, row_limit + 1):
                        row_vals = [f"**{r}**"]
                        for c in range(1, col_limit + 1):
                            cell_form = ws_form.cell(row=r, column=c).value
                            cell_val = ws_val.cell(row=r, column=c).value

                            if cell_form is not None and isinstance(cell_form, str) and cell_form.startswith("="):
                                if show_formulas:
                                    display = cell_form
                                else:
                                    val_str = str(cell_val) if cell_val is not None else ""
                                    display = f"{val_str} `({cell_form})`"
                            else:
                                display = str(cell_val) if cell_val is not None else ""

                            if len(display) > 40:
                                display = display[:37] + "..."
                            row_vals.append(display)
                        grid.append(row_vals)

                    sec.append(format_markdown_table(headers, grid))
                    if ws_form.max_row > max_rows and max_rows > 0:
                        sec.append(f"\n*Showing {row_limit} of {ws_form.max_row} rows and {col_limit} of {ws_form.max_column} columns.*")

                md_sections.append("\n".join(sec))

            output_result("\n\n---\n\n".join(md_sections), out_path)

    finally:
        wb_form.close()
        wb_val.close()


def inspect_xlsx_formulas(file_path: Path, target_sheet: Optional[str], fmt: str, out_path: Optional[Path]) -> None:
    """List all formula cells, coordinates, and expressions in Markdown, JSON, or HTML."""
    from openpyxl import load_workbook

    wb = load_workbook(file_path, data_only=False)
    try:
        sheets = [target_sheet] if target_sheet else wb.sheetnames
        formulas = []

        for sheet_name in sheets:
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            if not hasattr(ws, "iter_rows"):
                continue
            for row in ws.iter_rows():
                for cell in row:
                    val = cell.value
                    if isinstance(val, str) and val.startswith("="):
                        formulas.append({
                            "sheet": sheet_name,
                            "coordinate": cell.coordinate,
                            "formula": val,
                        })

        if fmt == "json":
            payload = {
                "file": file_path.name,
                "total_formulas": len(formulas),
                "formulas": formulas,
            }
            output_result(json.dumps(payload, indent=2, default=json_serializer, ensure_ascii=False), out_path)

        elif fmt == "html":
            table_lines = [
                f'<table class="excel-table"><caption>Formulas in: {html.escape(file_path.name)} ({len(formulas)} found)</caption>',
                '<thead><tr><th>Sheet</th><th>Coordinate</th><th>Formula</th></tr></thead>',
                '<tbody>'
            ]
            for f in formulas:
                table_lines.append(
                    f'<tr><td>{html.escape(f["sheet"])}</td>'
                    f'<td><strong>{html.escape(f["coordinate"])}</strong></td>'
                    f'<td><code>{html.escape(f["formula"])}</code></td></tr>'
                )
            table_lines.append('</tbody></table>')
            full_html = HTML_PAGE_TEMPLATE.format(
                title=f"Formulas: {file_path.name}",
                file_name=file_path.name,
                date=datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                content="\n".join(table_lines),
            )
            output_result(full_html, out_path)

        else:  # markdown
            headers = ["Sheet", "Coordinate", "Formula"]
            rows = [[f["sheet"], f["coordinate"], f["formula"]] for f in formulas]
            md_out = f"## Formulas in `{file_path.name}` ({len(formulas)} formula(s) found)\n\n"
            if rows:
                md_out += format_markdown_table(headers, rows)
            else:
                md_out += "*No formulas found.*"
            output_result(md_out, out_path)

    finally:
        wb.close()


def main():
    parser = argparse.ArgumentParser(description="Inspect structure, sheets, formulas, and convert spreadsheets to Markdown, JSON, or HTML.")
    parser.add_argument("file", type=Path, help="Path to .xlsx, .xlsm, .csv, or .tsv file")
    parser.add_argument(
        "--mode",
        choices=["summary", "preview", "formulas"],
        default="preview",
        help="Inspection mode: preview (tabular data, default), summary (sheet stats), formulas (all formulas)",
    )
    parser.add_argument(
        "--format",
        choices=["markdown", "json", "json-grid", "html"],
        default="markdown",
        help="Output format: markdown (default), json (records/summary), json-grid (cell-by-cell coordinate map), html (semantic table)",
    )
    parser.add_argument("--sheet", type=str, default=None, help="Target sheet name (default: active sheet)")
    parser.add_argument("--all-sheets", action="store_true", help="Process all sheets in workbook")
    parser.add_argument("--max-rows", type=int, default=50, help="Maximum rows to show in preview/export (default: 50, 0 for all)")
    parser.add_argument("--max-cols", type=int, default=20, help="Maximum cols to show in preview/export (default: 20, 0 for all)")
    parser.add_argument("--show-formulas", action="store_true", help="In preview, display raw formula instead of/with value")
    parser.add_argument("--dense", action="store_true", help="Omit null/empty cells in JSON output to save LLM tokens")
    parser.add_argument("--out", type=Path, default=None, help="Output file path (prints to stdout if not specified)")

    args = parser.parse_args()

    if not args.file.is_file():
        print(f"Error: File not found: {args.file}", file=sys.stderr)
        sys.exit(1)

    suffix = args.file.suffix.lower()
    if suffix in [".csv", ".tsv"]:
        inspect_csv(args.file, args.mode, args.format, args.max_rows, args.out)
    elif suffix in [".xlsx", ".xlsm", ".xltx"]:
        if args.mode == "summary":
            inspect_xlsx_summary(args.file, args.format, args.out)
        elif args.mode == "formulas":
            inspect_xlsx_formulas(args.file, args.sheet, args.format, args.out)
        else:  # preview / default
            inspect_xlsx_preview(
                args.file,
                args.sheet,
                args.all_sheets,
                args.max_rows,
                args.max_cols,
                args.show_formulas,
                args.format,
                args.dense,
                args.out,
            )
    else:
        print(f"Error: Unsupported file extension '{suffix}'. Expected .xlsx, .xlsm, .csv, or .tsv", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
